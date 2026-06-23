"""Master Data hub and reference-data tools."""
import csv
import io
import math
import os
import re
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP

from flask import Blueprint, Response, abort, flash, redirect, render_template, request, url_for

from app import config_store
from app.db import execute, get_db, insert_api_call_log, query_all, query_one
from app.hmrc_api import check_eori_batch
from app.pipeline_validation import normalise_package_type
from app.tenant import TENANT_REGISTRY, get_tenant


master_data_bp = Blueprint("master_data", __name__, template_folder="../../templates/master_data")

PARTNER_TYPES = ["Carrier", "Haulier", "Consignor", "Consignee", "Importer", "Exporter"]
STATIC_CV_TABLE_NAMES = [
    "CV_addition_deduction_code",
    "CV_additional_info_code",
    "CV_additional_procedure_code",
    "CV_ap_auth_role_code",
    "CV_ap_auth_role_type",
    "CV_auth_role_type",
    "CV_auth_type_code",
    "CV_cargo_or_consignment",
    "CV_commodity_code",
    "CV_controlled_goods_type",
    "CV_country",
    "CV_currency",
    "CV_declaration_category",
    "CV_document_code",
    "CV_document_status",
    "CV_ffd_declaration_choice",
    "CV_ffd_location_of_goods",
    "CV_final_destination_location_code",
    "CV_goods_domestic_status",
    "CV_guarantee_type",
    "CV_gvms_routes",
    "CV_incoterm",
    "CV_inland_mode_of_transport",
    "CV_item_add_ded_code",
    "CV_load_type",
    "CV_measurement_unit",
    "CV_method_of_payment",
    "CV_mode_of_transport",
    "CV_movement_type",
    "CV_national_additional_code",
    "CV_nature_of_transaction",
    "CV_ni_additional_information_code",
    "CV_no_sfd_reason",
    "CV_passive_transport_types",
    "CV_port",
    "CV_port_stg",
    "CV_preference",
    "CV_previous_document_class",
    "CV_previous_document_type",
    "CV_procedure_code",
    "CV_representation_type",
    "CV_route",
    "CV_sd_declaration_choice",
    "CV_sd_location_of_goods",
    "CV_sd_status",
    "CV_sfd_declaration_choice",
    "CV_sfd_header_movement_type",
    "CV_special_authorisation",
    "CV_standalone_sdi_authorisation_type",
    "CV_supervising_customs_office",
    "CV_tax_base_unit",
    "CV_tax_type",
    "CV_transport_charge",
    "CV_transport_document_type",
    "CV_type_of_package",
    "CV_un_locode",
    "CV_valuation_indicator",
    "CV_valuation_method",
]
CV_PAGE_SIZE = 50
VALIDATION_SETTINGS = [
    {
        "category": "VALIDATION",
        "key": "STRICT_MASTERDATA_VALIDATION",
        "label": "Strict local masterdata validation",
        "fallback": "false",
        "help": (
            "When enabled, local validation blocks unknown/mismatched party EORIs, placeholder EORIs "
            "and missing ENS carrier masterdata before sending to TSS. Keep disabled for end-to-end "
            "test runs with random values."
        ),
    },
]

PRODUCT_EXPORT_FIELDS = [
    "source_table",
    "id",
    "customer_code",
    "sku",
    "product_code",
    "barcode",
    "product_name",
    "goods_description",
    "commodity_code",
    "country_of_origin",
    "package_type",
    "package_marks",
    "procedure_code",
    "additional_procedure_code",
    "valuation_method",
    "valuation_indicator",
    "preference_code",
    "ni_additional_info_code",
    "nature_of_transaction",
    "country_of_preferential_origin",
    "taric_code",
    "cus_code",
    "national_additional_code",
    "quota_order_number",
    "controlled_goods_type",
    "sdi_notes",
    "gross_weight_kg",
    "net_weight_kg",
    "weight_source",
    "weight_sample_count",
    "unit_value",
    "currency",
    "statistical_unit",
    "controlled_goods",
    "requires_supplementary_unit",
    "is_active",
    "notes",
    "created_at",
    "updated_at",
]

PRODUCT_IMPORT_ALIASES = {
    "customer_code": ("customer_code", "customer", "customer code"),
    "sku": ("sku", "item", "item no", "item_no", "no", "product sku", "code sku", "code / sku"),
    "product_code": ("product_code", "product code", "code", "stock_code", "stock code", "code sku", "code / sku"),
    "barcode": ("barcode", "bar code", "ean"),
    "description": ("description", "product_name", "product name", "goods_description", "goods description"),
    "commodity_code": ("commodity_code", "commodity code", "hs_code", "hs code", "tariff code"),
    "taric_code": ("taric_code", "taric code", "taric"),
    "cus_code": ("cus_code", "cus code", "cus"),
    "national_additional_code": ("national_additional_code", "national additional code", "national add code", "nac"),
    "quota_order_number": ("quota_order_number", "quota order number", "quota", "quota order no"),
    "country_of_origin": (
        "country_of_origin",
        "country of origin",
        "origin",
        "country",
        "item origin",
        "country region of origin code",
    ),
    "country_of_preferential_origin": (
        "country_of_preferential_origin",
        "country of preferential origin",
        "preferential origin",
        "preference origin",
    ),
    "gross_weight_kg": ("gross_weight_kg", "gross weight kg", "gross_mass_kg", "gross mass kg", "gross kg"),
    "net_weight_kg": ("net_weight_kg", "net weight kg", "net_mass_kg", "net mass kg", "net kg"),
    "package_type": (
        "package_type",
        "package type",
        "type_of_packages",
        "type of packages",
        "unit of measure code",
        "uom",
        "uom code",
    ),
    "procedure_code": ("procedure_code", "procedure code", "procedure", "pc"),
    "additional_procedure_code": (
        "additional_procedure_code",
        "additional procedure code",
        "additional procedure",
        "apc",
    ),
    "valuation_method": ("valuation_method", "valuation method"),
    "valuation_indicator": ("valuation_indicator", "valuation indicator"),
    "preference_code": ("preference_code", "preference", "preference code"),
    "ni_additional_information_codes": (
        "ni_additional_information_codes",
        "ni_additional_info_code",
        "ni additional information codes",
        "ni additional info",
        "ni ai",
    ),
    "nature_of_transaction": ("nature_of_transaction", "nature of transaction", "transaction nature", "not"),
    "unit_price": ("unit_price", "unit value", "unit_value", "unit_value_gbp", "unit price"),
    "currency": ("currency", "item_invoice_currency", "invoice currency"),
    "statistical_unit": ("statistical_unit", "statistical unit"),
    "controlled_goods": ("controlled_goods", "controlled goods", "controlled"),
    "controlled_goods_type": ("controlled_goods_type", "controlled goods type", "controlled type"),
    "sdi_notes": ("sdi_notes", "sdi notes", "supdec notes", "supplementary declaration notes"),
    "requires_supplementary_unit": ("requires_supplementary_unit", "requires supplementary unit", "supplementary unit"),
    "active": ("active", "is_active", "status"),
    "weight_source": ("weight_source", "weight source"),
    "weight_sample_count": ("weight_sample_count", "weight sample count", "samples"),
}
PRODUCT_IMPORT_BACKFILL_FIELDS = (
    "description",
    "commodity_code",
    "country_of_origin",
    "package_type",
    "procedure_code",
    "additional_procedure_code",
    "preference_code",
    "ni_additional_information_codes",
    "ni_additional_info_code",
    "nature_of_transaction",
    "valuation_method",
    "valuation_indicator",
    "country_of_preferential_origin",
    "taric_code",
    "cus_code",
    "national_additional_code",
    "quota_order_number",
    "controlled_goods_type",
    "gross_weight_kg",
    "net_weight_kg",
    "unit_price",
    "currency",
    "statistical_unit",
    "weight_source",
    "weight_sample_count",
)
PRODUCT_INFO_STATUS_OPTIONS = [
    ("", "All products"),
    ("complete", "Complete information"),
    ("incomplete", "Missing information"),
    ("missing_weights", "Missing weights"),
]
PRODUCT_INFO_STATUS_VALUES = {value for value, _label in PRODUCT_INFO_STATUS_OPTIONS}
PRODUCT_WEIGHT_FETCH_SOURCE = "Historical goods fetch"
PRODUCT_WEIGHT_LOOKUP_FIELDS = ("sku", "product_code", "stock_code", "barcode")
PRODUCT_WEIGHT_FETCH_ROW_LIMIT = 500
PRODUCT_WEIGHT_FETCH_SAMPLE_LIMIT = 10000
INGEST_DEFAULT_SECTIONS = [
    {
        "title": "Ingestion Flow",
        "icon": "bi-sliders",
        "intro": "Tenant-level defaults for local invoice intake. If the app says record creation is disabled, set Ingestion creation enabled to Enabled and save this page.",
        "fields": [
            {
                "key": "ENABLED",
                "label": "Ingestion creation enabled",
                "anchor": "ingestCreationEnabled",
                "type": "select",
                "options": [("true", "Enabled"), ("false", "Disabled")],
                "fallback": "true",
                "help": "This is the switch from the warning toast. Choose Enabled, then Save Company, to let Confirm create local ENS, consignments and goods. It does not submit records to TSS.",
            },
            {
                "key": "AUTO_VALIDATE",
                "label": "Auto-validate after staging",
                "type": "select",
                "options": [("true", "Enabled"), ("false", "Disabled")],
                "fallback": "true",
                "help": "Runs local validation after ENS, consignments and goods have been staged. This does not unlock creation by itself.",
            },
            {
                "key": "SUPPLIER_NAME",
                "label": "Supplier Name",
                "type": "text",
                "placeholder": "Supplier / exporter name used for matching",
                "fallback": "",
                "help": "Used to resolve consignor/exporter partner defaults during invoice ingest.",
            },
            {
                "key": "DEFAULT_ARRIVAL_HOURS_AHEAD",
                "label": "Arrival Hours Ahead",
                "type": "number",
                "placeholder": "4",
                "fallback": "4",
                "help": "If the inbound document does not provide an arrival timestamp, use now plus this many hours.",
            },
        ],
    },
    {
        "title": "Transport & Carrier Defaults",
        "icon": "bi-truck",
        "intro": "Operational defaults reused when the tenant stages ENS headers from inbound documents.",
        "fields": [
            {"key": "DEFAULT_MOVEMENT_TYPE", "label": "Movement Type", "type": "text", "placeholder": "1", "fallback": "1"},
            {"key": "DEFAULT_IDENTITY_NO_OF_TRANSPORT", "label": "Transport Identity", "type": "text", "placeholder": "IMO1234567", "fallback": "IMO1234567"},
            {"key": "DEFAULT_NATIONALITY_OF_TRANSPORT", "label": "Transport Nationality", "type": "text", "placeholder": "GB", "fallback": "GB"},
            {"key": "DEFAULT_ARRIVAL_PORT", "label": "Arrival Port", "type": "text", "placeholder": "DUBLIN", "fallback": ""},
            {"key": "DEFAULT_PLACE_OF_LOADING", "label": "Place of Loading", "type": "text", "placeholder": "FRPAR", "fallback": ""},
            {"key": "DEFAULT_PLACE_OF_UNLOADING", "label": "Place of Unloading", "type": "text", "placeholder": "DUBLIN", "fallback": ""},
            {"key": "DEFAULT_TRANSPORT_CHARGES", "label": "Transport Charges", "type": "text", "placeholder": "A", "fallback": "A"},
            {"key": "DEFAULT_CARRIER_EORI", "label": "Carrier EORI", "type": "text", "placeholder": "XI123456789000", "fallback": ""},
            {"key": "DEFAULT_CARRIER_NAME", "label": "Carrier Name", "type": "text", "placeholder": "Carrier name", "fallback": ""},
            {"key": "DEFAULT_CARRIER_STREET_NUMBER", "label": "Carrier Street / Number", "type": "text", "placeholder": "10 Harbour Road", "fallback": ""},
            {"key": "DEFAULT_CARRIER_CITY", "label": "Carrier City", "type": "text", "placeholder": "Belfast", "fallback": ""},
            {"key": "DEFAULT_CARRIER_POSTCODE", "label": "Carrier Postcode", "type": "text", "placeholder": "BT1 1AA", "fallback": ""},
            {"key": "DEFAULT_CARRIER_COUNTRY", "label": "Carrier Country", "type": "text", "placeholder": "GB", "fallback": "GB"},
            {"key": "DEFAULT_HAULIER_EORI", "label": "Haulier EORI", "type": "text", "placeholder": "GB123456789000", "fallback": ""},
        ],
    },
    {
        "title": "Consignment & Goods Defaults",
        "icon": "bi-box-seam",
        "intro": "Fallback values used when invoice lines do not provide enough customs detail.",
        "fields": [
            {
                "key": "DEFAULT_CONTAINER_INDICATOR",
                "label": "Container Indicator",
                "type": "select",
                "options": [("0", "0 - Uncontainerised"), ("1", "1 - Containerised")],
                "fallback": "0",
            },
            {
                "key": "DEFAULT_CONTROLLED_GOODS",
                "label": "Controlled Goods",
                "type": "select",
                "options": [("yes", "Yes"), ("no", "No")],
                "fallback": "no",
            },
            {"key": "DEFAULT_COUNTRY_OF_ORIGIN", "label": "Country of Origin", "type": "text", "placeholder": "GB", "fallback": "GB"},
            {"key": "DEFAULT_PACKAGE_TYPE", "label": "Package Type", "type": "text", "placeholder": "PK", "fallback": "PK"},
            {"key": "DEFAULT_PROCEDURE_CODE", "label": "Procedure Code", "type": "text", "placeholder": "4000", "fallback": "4000"},
            {"key": "DEFAULT_ADDITIONAL_PROCEDURE_CODE", "label": "Additional Procedure", "type": "text", "placeholder": "000", "fallback": "000"},
            {"key": "DEFAULT_VALUATION_METHOD", "label": "Valuation Method", "type": "text", "placeholder": "1", "fallback": "1"},
            {"key": "DEFAULT_INVOICE_CURRENCY", "label": "Invoice Currency", "type": "text", "placeholder": "GBP", "fallback": "GBP"},
            {
                "key": "DEFAULT_IMPORTER_EORI",
                "label": "SFD Importer EORI",
                "type": "text",
                "placeholder": "XI123456789000",
                "fallback": "",
                "help": "When Use SFD is selected, ingestion can use this TSS-registered importer instead of a customer EORI that may force no_sfd_reason. If blank, Company Master EORI is used when available.",
            },
            {
                "key": "DEFAULT_IMPORTER_NAME",
                "label": "SFD Importer Name",
                "type": "text",
                "placeholder": "Importer legal name",
                "fallback": "",
                "help": "Fallback importer name sent with the SFD importer EORI when TSS needs address data.",
            },
            {
                "key": "DEFAULT_IMPORTER_STREET_NUMBER",
                "label": "SFD Importer Street / Number",
                "type": "text",
                "placeholder": "10 Import Road",
                "fallback": "",
            },
            {"key": "DEFAULT_IMPORTER_CITY", "label": "SFD Importer City", "type": "text", "placeholder": "Belfast", "fallback": ""},
            {"key": "DEFAULT_IMPORTER_POSTCODE", "label": "SFD Importer Postcode", "type": "text", "placeholder": "BT1 1AA", "fallback": ""},
            {"key": "DEFAULT_IMPORTER_COUNTRY", "label": "SFD Importer Country", "type": "text", "placeholder": "GB", "fallback": "GB"},
            {"key": "DEFAULT_CONSIGNOR_EORI", "label": "Consignor EORI", "type": "text", "placeholder": "GB123456789000", "fallback": ""},
            {"key": "DEFAULT_EXPORTER_EORI", "label": "Exporter EORI", "type": "text", "placeholder": "GB123456789000", "fallback": ""},
        ],
    },
]


def _table_exists(schema_name, table_name):
    row = query_one(
        """
        SELECT 1 AS ok
        FROM INFORMATION_SCHEMA.TABLES
        WHERE TABLE_SCHEMA = ? AND TABLE_NAME = ?
        """,
        [schema_name, table_name],
    )
    return bool(row)


def _table_columns(schema_name, table_name):
    try:
        rows = query_all(
            """
            SELECT COLUMN_NAME
            FROM INFORMATION_SCHEMA.COLUMNS
            WHERE TABLE_SCHEMA = ? AND TABLE_NAME = ?
            """,
            [schema_name, table_name],
        )
        return {row.get("COLUMN_NAME") for row in rows if row.get("COLUMN_NAME")}
    except Exception:
        return set()


def _tenant_ctx():
    tenant = get_tenant()
    return tenant, tenant["schema"]


def _product_target_ctx():
    tenant = get_tenant()
    requested = (
        request.values.get("tenant_code")
        or request.values.get("product_tenant")
        or ""
    ).strip().upper()
    if requested and requested != "ALL":
        if tenant.get("code") == "SYD" or tenant.get("code") == requested:
            target = TENANT_REGISTRY.get(requested)
            if target:
                tenant = target
    return tenant, tenant["schema"]


def _products_url_for_tenant(tenant, **kwargs):
    if tenant and tenant.get("code"):
        kwargs.setdefault("product_tenant", tenant["code"])
    return url_for("master_data.products", **kwargs)


def _qualified(schema_name, table_name):
    return f"[{schema_name}].[{table_name}]"


def _tenant_config_field_map():
    field_map = {}
    for field in VALIDATION_SETTINGS:
        field_map[field["key"]] = field
    for section in INGEST_DEFAULT_SECTIONS:
        for field in section["fields"]:
            field_map[field["key"]] = field
    return field_map


def _company_record(schema_name=None):
    schema_name = schema_name or _tenant_ctx()[1]
    try:
        return query_one(f"SELECT TOP 1 * FROM {_qualified(schema_name, 'CompanyMaster')} ORDER BY id")
    except Exception:
        return None


def _safe_count(sql, params=None):
    try:
        row = query_one(sql, params or [])
        return int((row or {}).get("cnt", 0))
    except Exception:
        return None


def _product_active_column(columns):
    if "is_active" in columns:
        return "is_active"
    if "active" in columns:
        return "active"
    return None


def _select_existing_expr(columns, alias, candidates, fallback_sql):
    for candidate in candidates:
        if candidate in columns:
            return f"{candidate} AS {alias}"
    return f"{fallback_sql} AS {alias}"


def _coalesce_text_expr(columns, alias, candidates, fallback_sql="CAST(NULL AS NVARCHAR(200))"):
    present = [f"NULLIF({candidate}, '')" for candidate in candidates if candidate in columns]
    if len(present) == 1:
        return f"{present[0]} AS {alias}"
    if present:
        return f"COALESCE({', '.join(present)}) AS {alias}"
    return f"{fallback_sql} AS {alias}"


def _product_select_sql(schema_name, columns):
    active_col = _product_active_column(columns)
    active_expr = active_col if active_col else "CAST(1 AS bit)"
    select_expr = ", ".join([
        _select_existing_expr(columns, "id", ("id",), "CAST(NULL AS int)"),
        _coalesce_text_expr(columns, "product_code", ("product_code", "sku", "barcode")),
        _coalesce_text_expr(columns, "product_name", ("product_name", "goods_description", "description")),
        _coalesce_text_expr(columns, "goods_description", ("goods_description", "product_name", "description")),
        _select_existing_expr(columns, "commodity_code", ("commodity_code",), "CAST(NULL AS NVARCHAR(20))"),
        _select_existing_expr(columns, "country_of_origin", ("country_of_origin",), "CAST(NULL AS NVARCHAR(3))"),
        _select_existing_expr(columns, "package_type", ("package_type", "type_of_packages"), "CAST(NULL AS NVARCHAR(10))"),
        _select_existing_expr(columns, "package_marks", ("package_marks",), "CAST(NULL AS NVARCHAR(256))"),
        _select_existing_expr(columns, "procedure_code", ("procedure_code",), "CAST(NULL AS NVARCHAR(10))"),
        _select_existing_expr(columns, "additional_procedure_code", ("additional_procedure_code",), "CAST(NULL AS NVARCHAR(10))"),
        _select_existing_expr(columns, "valuation_method", ("valuation_method",), "CAST(NULL AS NVARCHAR(5))"),
        _select_existing_expr(columns, "valuation_indicator", ("valuation_indicator",), "CAST(NULL AS NVARCHAR(20))"),
        _select_existing_expr(columns, "preference_code", ("preference_code", "preference"), "CAST(NULL AS NVARCHAR(10))"),
        _select_existing_expr(columns, "ni_additional_info_code", ("ni_additional_info_code", "ni_additional_information_codes"), "CAST(NULL AS NVARCHAR(40))"),
        _select_existing_expr(columns, "nature_of_transaction", ("nature_of_transaction",), "CAST(NULL AS NVARCHAR(40))"),
        _select_existing_expr(columns, "country_of_preferential_origin", ("country_of_preferential_origin",), "CAST(NULL AS NVARCHAR(2))"),
        _select_existing_expr(columns, "taric_code", ("taric_code",), "CAST(NULL AS NVARCHAR(20))"),
        _select_existing_expr(columns, "cus_code", ("cus_code",), "CAST(NULL AS NVARCHAR(20))"),
        _select_existing_expr(columns, "national_additional_code", ("national_additional_code",), "CAST(NULL AS NVARCHAR(10))"),
        _select_existing_expr(columns, "quota_order_number", ("quota_order_number",), "CAST(NULL AS NVARCHAR(10))"),
        _select_existing_expr(columns, "controlled_goods_type", ("controlled_goods_type",), "CAST(NULL AS NVARCHAR(40))"),
        _select_existing_expr(columns, "gross_weight_kg", ("default_gross_weight_kg", "gross_weight_kg", "gross_mass_kg"), "CAST(NULL AS DECIMAL(18,3))"),
        _select_existing_expr(columns, "net_weight_kg", ("default_net_weight_kg", "net_weight_kg", "net_mass_kg"), "CAST(NULL AS DECIMAL(18,3))"),
        _select_existing_expr(columns, "unit_value", ("unit_value_gbp", "unit_price", "unit_value"), "CAST(NULL AS DECIMAL(18,4))"),
        _select_existing_expr(columns, "currency", ("currency",), "CAST('GBP' AS CHAR(3))"),
        _select_existing_expr(columns, "statistical_unit", ("statistical_unit",), "CAST(NULL AS NVARCHAR(20))"),
        _select_existing_expr(columns, "controlled_goods", ("controlled_goods",), "CAST(0 AS bit)"),
        _select_existing_expr(columns, "requires_supplementary_unit", ("requires_supplementary_unit",), "CAST(0 AS bit)"),
        _select_existing_expr(columns, "weight_source", ("weight_source",), "CAST(NULL AS NVARCHAR(160))"),
        _select_existing_expr(columns, "weight_sample_count", ("weight_sample_count",), "CAST(NULL AS int)"),
        f"{active_expr} AS is_active",
        _select_existing_expr(columns, "notes", ("notes",), "CAST(NULL AS NVARCHAR(500))"),
        _select_existing_expr(columns, "sdi_notes", ("sdi_notes",), "CAST(NULL AS NVARCHAR(1000))"),
        _select_existing_expr(columns, "created_at", ("created_at",), "CAST(NULL AS DATETIME2)"),
        _select_existing_expr(columns, "updated_at", ("updated_at",), "CAST(NULL AS DATETIME2)"),
        "CAST(NULL AS NVARCHAR(20)) AS customer_code",
        "CAST(NULL AS NVARCHAR(100)) AS sku",
        "CAST(NULL AS NVARCHAR(100)) AS barcode",
        "'Products' AS source_table",
    ])
    return f"SELECT {select_expr} FROM {_qualified(schema_name, 'Products')}"


def _product_order_clause(columns):
    order_cols = []
    active_col = _product_active_column(columns)
    if active_col:
        order_cols.append(f"{active_col} DESC")
    if "product_name" in columns:
        order_cols.append("product_name")
    if "product_code" in columns:
        order_cols.append("product_code")
    if not order_cols:
        order_cols.append("id")
    return " ORDER BY " + ", ".join(order_cols)


def _product_search_clause(columns):
    searchable = [
        col
        for col in ("product_name", "product_code", "goods_description", "commodity_code", "country_of_origin")
        if col in columns
    ]
    if not searchable:
        return "", []
    return "(" + " OR ".join(f"{col} LIKE ?" for col in searchable) + ")", searchable


def _normalize_product_info_status(value):
    value = (value or "").strip().lower()
    return value if value in PRODUCT_INFO_STATUS_VALUES else ""


def _has_product_text(row, *keys):
    return any(str(row.get(key) or "").strip() for key in keys)


def _has_positive_product_number(row, *keys):
    for key in keys:
        value = row.get(key)
        if value in (None, ""):
            continue
        try:
            if float(value) > 0:
                return True
        except (TypeError, ValueError):
            continue
    return False


def _product_missing_info_fields(row):
    missing = []
    if not _has_product_text(row, "product_name", "goods_description"):
        missing.append("description")
    if not _has_product_text(row, "commodity_code"):
        missing.append("commodity code")
    if not _has_product_text(row, "country_of_origin"):
        missing.append("origin")
    if not _has_positive_product_number(row, "gross_weight_kg"):
        missing.append("gross weight")
    if not _has_positive_product_number(row, "net_weight_kg"):
        missing.append("net weight")
    return missing


def _annotate_product_info_status(rows):
    annotated = []
    for row in rows:
        item = dict(row)
        missing = _product_missing_info_fields(item)
        item["missing_info_fields"] = missing
        item["info_complete"] = not missing
        item["missing_weights"] = any(field in missing for field in ("gross weight", "net weight"))
        annotated.append(item)
    return annotated


def _filter_product_info_status(rows, info_status):
    info_status = _normalize_product_info_status(info_status)
    annotated = _annotate_product_info_status(rows)
    if info_status == "complete":
        return [row for row in annotated if row["info_complete"]]
    if info_status == "incomplete":
        return [row for row in annotated if not row["info_complete"]]
    if info_status == "missing_weights":
        return [row for row in annotated if row["missing_weights"]]
    return annotated


def _doc_product_catalog_count(schema_name):
    if not _table_exists(schema_name, "DocProductCatalog"):
        return None
    columns = _table_columns(schema_name, "DocProductCatalog")
    where = " WHERE active = 1" if "active" in columns else ""
    return _safe_count(f"SELECT COUNT(*) AS cnt FROM {_qualified(schema_name, 'DocProductCatalog')}{where}")


def _doc_product_catalog_rows(schema_name, search, *, active_only=True):
    if not _table_exists(schema_name, "DocProductCatalog"):
        return []
    columns = _table_columns(schema_name, "DocProductCatalog")
    active_where = "active = 1" if active_only and "active" in columns else ""
    searchable = [
        col
        for col in ("sku", "barcode", "product_code", "description", "commodity_code", "customer_code", "country_of_origin")
        if col in columns
    ]
    where = []
    params = []
    if active_where:
        where.append(active_where)
    if search and searchable:
        where.append("(" + " OR ".join(f"{col} LIKE ?" for col in searchable) + ")")
        params.extend([f"%{search}%" for _ in searchable])
    select_expr = ", ".join([
        _select_existing_expr(columns, "id", ("id",), "CAST(NULL AS int)"),
        _coalesce_text_expr(columns, "product_code", ("product_code", "sku", "barcode")),
        _coalesce_text_expr(columns, "product_name", ("description", "product_code", "sku")),
        _coalesce_text_expr(columns, "goods_description", ("description", "product_code", "sku")),
        _select_existing_expr(columns, "commodity_code", ("commodity_code",), "CAST(NULL AS NVARCHAR(20))"),
        _select_existing_expr(columns, "country_of_origin", ("country_of_origin",), "CAST(NULL AS NVARCHAR(3))"),
        _select_existing_expr(columns, "package_type", ("package_type", "type_of_packages"), "CAST(NULL AS NVARCHAR(40))"),
        "CAST(NULL AS NVARCHAR(256)) AS package_marks",
        _select_existing_expr(columns, "procedure_code", ("procedure_code",), "CAST(NULL AS NVARCHAR(10))"),
        _select_existing_expr(columns, "additional_procedure_code", ("additional_procedure_code",), "CAST(NULL AS NVARCHAR(10))"),
        _select_existing_expr(columns, "valuation_method", ("valuation_method",), "CAST(NULL AS NVARCHAR(5))"),
        _select_existing_expr(columns, "valuation_indicator", ("valuation_indicator",), "CAST(NULL AS NVARCHAR(20))"),
        _select_existing_expr(columns, "preference_code", ("preference_code", "preference"), "CAST(NULL AS NVARCHAR(10))"),
        _select_existing_expr(columns, "ni_additional_info_code", ("ni_additional_info_code", "ni_additional_information_codes"), "CAST(NULL AS NVARCHAR(40))"),
        _select_existing_expr(columns, "nature_of_transaction", ("nature_of_transaction",), "CAST(NULL AS NVARCHAR(40))"),
        _select_existing_expr(columns, "country_of_preferential_origin", ("country_of_preferential_origin",), "CAST(NULL AS NVARCHAR(2))"),
        _select_existing_expr(columns, "taric_code", ("taric_code",), "CAST(NULL AS NVARCHAR(20))"),
        _select_existing_expr(columns, "cus_code", ("cus_code",), "CAST(NULL AS NVARCHAR(20))"),
        _select_existing_expr(columns, "national_additional_code", ("national_additional_code",), "CAST(NULL AS NVARCHAR(10))"),
        _select_existing_expr(columns, "quota_order_number", ("quota_order_number",), "CAST(NULL AS NVARCHAR(10))"),
        _select_existing_expr(columns, "controlled_goods_type", ("controlled_goods_type",), "CAST(NULL AS NVARCHAR(40))"),
        _select_existing_expr(columns, "gross_weight_kg", ("gross_weight_kg", "default_gross_weight_kg", "gross_mass_kg"), "CAST(NULL AS DECIMAL(18,3))"),
        _select_existing_expr(columns, "net_weight_kg", ("net_weight_kg", "default_net_weight_kg", "net_mass_kg"), "CAST(NULL AS DECIMAL(18,3))"),
        _select_existing_expr(columns, "unit_value", ("unit_price", "unit_value_gbp", "unit_value"), "CAST(NULL AS DECIMAL(18,4))"),
        _select_existing_expr(columns, "currency", ("currency",), "CAST(NULL AS CHAR(3))"),
        _select_existing_expr(columns, "statistical_unit", ("statistical_unit",), "CAST(NULL AS NVARCHAR(20))"),
        _select_existing_expr(columns, "controlled_goods", ("controlled_goods",), "CAST(0 AS bit)"),
        _select_existing_expr(columns, "requires_supplementary_unit", ("requires_supplementary_unit",), "CAST(0 AS bit)"),
        _select_existing_expr(columns, "weight_source", ("weight_source",), "CAST(NULL AS NVARCHAR(160))"),
        _select_existing_expr(columns, "weight_sample_count", ("weight_sample_count",), "CAST(NULL AS int)"),
        f"{'active' if 'active' in columns else 'CAST(1 AS bit)'} AS is_active",
        _select_existing_expr(columns, "notes", ("notes", "sdi_notes"), "CAST(NULL AS NVARCHAR(500))"),
        _select_existing_expr(columns, "sdi_notes", ("sdi_notes",), "CAST(NULL AS NVARCHAR(1000))"),
        _select_existing_expr(columns, "created_at", ("created_at",), "CAST(NULL AS DATETIME2)"),
        _select_existing_expr(columns, "updated_at", ("updated_at",), "CAST(NULL AS DATETIME2)"),
        _select_existing_expr(columns, "customer_code", ("customer_code",), "CAST(NULL AS NVARCHAR(20))"),
        _select_existing_expr(columns, "sku", ("sku",), "CAST(NULL AS NVARCHAR(100))"),
        _select_existing_expr(columns, "barcode", ("barcode",), "CAST(NULL AS NVARCHAR(100))"),
        "'DocProductCatalog' AS source_table",
    ])
    sql = f"SELECT {select_expr} FROM {_qualified(schema_name, 'DocProductCatalog')}"
    if where:
        sql += " WHERE " + " AND ".join(where)
    sql += " ORDER BY product_code, customer_code"
    return query_all(sql, params)


def _load_product_master_rows(schema_name, search="", *, catalog_active_only=True, info_status=""):
    has_products_table = _table_exists(schema_name, "Products")
    has_catalog_table = _table_exists(schema_name, "DocProductCatalog")
    rows = []
    product_source_label = "Products"

    if has_products_table:
        columns = _table_columns(schema_name, "Products")
        sql = _product_select_sql(schema_name, columns)
        params = []
        if search:
            like = f"%{search}%"
            search_clause, searchable = _product_search_clause(columns)
            if search_clause:
                sql += " WHERE " + search_clause
                params = [like for _ in searchable]
        sql += _product_order_clause(columns)
        rows = query_all(sql, params)

    if has_catalog_table:
        catalog_rows = _doc_product_catalog_rows(schema_name, search, active_only=catalog_active_only)
        if rows and catalog_rows:
            rows.extend(catalog_rows)
            product_source_label = "Products + DocProductCatalog"
        elif catalog_rows:
            rows = catalog_rows
            product_source_label = "DocProductCatalog"

    rows = _filter_product_info_status(rows, info_status)

    return {
        "rows": rows,
        "has_products_table": has_products_table,
        "has_catalog_table": has_catalog_table,
        "has_product_source": has_products_table or has_catalog_table,
        "product_source_label": product_source_label,
    }


def _product_tenant_options(active_tenant):
    if (active_tenant or {}).get("code") != "SYD":
        return []
    options = [{"code": "ALL", "name": "All tenants"}]
    options.extend(
        {"code": item["code"], "name": item["name"], "schema": item["schema"]}
        for item in sorted(TENANT_REGISTRY.values(), key=lambda row: row["code"])
    )
    return options


def _selected_product_tenant(active_tenant, raw_value):
    if (active_tenant or {}).get("code") != "SYD":
        return (active_tenant or {}).get("code") or "BKD"
    requested = str(raw_value or active_tenant.get("code") or "SYD").strip().upper()
    if requested == "ALL":
        return "ALL"
    return requested if requested in TENANT_REGISTRY else active_tenant.get("code", "SYD")


def _product_scope(active_tenant, raw_value):
    selected_product_tenant = _selected_product_tenant(active_tenant, raw_value)
    if selected_product_tenant == "ALL":
        tenants = [
            item for item in sorted(TENANT_REGISTRY.values(), key=lambda row: row["code"])
            if item.get("schema")
        ]
        return selected_product_tenant, tenants, "All tenants"

    scope_tenant = TENANT_REGISTRY.get(selected_product_tenant, active_tenant)
    return selected_product_tenant, [scope_tenant], f"{scope_tenant['code']} - {scope_tenant['name']}"


def _annotate_product_rows_for_tenant(rows, tenant):
    annotated = []
    for row in rows:
        item = dict(row)
        item["tenant_code"] = tenant["code"]
        item["tenant_name"] = tenant["name"]
        item["tenant_schema"] = tenant["schema"]
        annotated.append(item)
    return annotated


def _load_product_master_rows_for_tenants(tenants, search="", *, catalog_active_only=True, info_status=""):
    combined_rows = []
    has_products_table = False
    has_catalog_table = False
    source_labels = set()

    for tenant in tenants:
        loaded = _load_product_master_rows(
            tenant["schema"],
            search,
            catalog_active_only=catalog_active_only,
            info_status=info_status,
        )
        has_products_table = has_products_table or loaded["has_products_table"]
        has_catalog_table = has_catalog_table or loaded["has_catalog_table"]
        if loaded["has_product_source"]:
            source_labels.add(loaded["product_source_label"])
        combined_rows.extend(_annotate_product_rows_for_tenant(loaded["rows"], tenant))

    combined_rows.sort(key=lambda row: (
        str(row.get("tenant_code") or ""),
        str(row.get("product_code") or row.get("sku") or row.get("barcode") or ""),
        str(row.get("customer_code") or ""),
    ))

    if not source_labels:
        product_source_label = "Products"
    elif len(source_labels) == 1:
        product_source_label = next(iter(source_labels))
    else:
        product_source_label = "Mixed product sources"

    return {
        "rows": combined_rows,
        "has_products_table": has_products_table,
        "has_catalog_table": has_catalog_table,
        "has_product_source": has_products_table or has_catalog_table,
        "product_source_label": product_source_label,
    }


def _doc_product_catalog_count_for_tenants(tenants):
    total = 0
    found = False
    for tenant in tenants:
        count = _doc_product_catalog_count(tenant["schema"])
        if count is not None:
            total += count
            found = True
    return total if found else None


def _raw_data_counts():
    tables = {
        "customers": "BKD_Customers",
        "products": "BKD_Item_Commodity_Code",
        "sales_orders": "BKD_Sales_Orders",
        "sales_orders_daily": "BKD_Sales_Orders_Daily",
    }
    counts = {}
    for key, table in tables.items():
        counts[key] = (
            _safe_count(f"SELECT COUNT(*) AS cnt FROM {_qualified('DATA', table)}")
            if _table_exists("DATA", table)
            else None
        )
    return counts


def _eori_cache_stats(schema_name):
    if not _table_exists(schema_name, "PrecheckEoriCache"):
        return None
    try:
        return query_one(
            f"""
            SELECT COUNT(*) AS total,
                   SUM(CASE WHEN is_valid = 1 THEN 1 ELSE 0 END) AS valid,
                   SUM(CASE WHEN is_valid = 0 THEN 1 ELSE 0 END) AS invalid
            FROM {_qualified(schema_name, 'PrecheckEoriCache')}
            """
        )
    except Exception:
        return None


def _product_form_payload(form, columns):
    values = {
        "product_code": form.get("product_code"),
        "product_name": form.get("product_name"),
        "commodity_code": form.get("commodity_code") or None,
        "goods_description": form.get("goods_description") or None,
        "country_of_origin": form.get("country_of_origin") or None,
        "package_type": form.get("package_type") or None,
        "package_marks": form.get("package_marks") or "No marks",
        "procedure_code": form.get("procedure_code") or "4000",
        "default_gross_weight_kg": _to_float(form.get("default_gross_weight_kg")),
        "default_net_weight_kg": _to_float(form.get("default_net_weight_kg")),
        "unit_value_gbp": _to_float(form.get("unit_value_gbp")),
        "valuation_method": form.get("valuation_method") or "1",
        "ni_additional_info_code": form.get("ni_additional_info_code") or None,
        "preference_code": form.get("preference_code") or "100",
        "notes": form.get("notes") or None,
    }
    active_col = _product_active_column(columns)
    if active_col:
        values[active_col] = 1 if form.get("is_active") else 0
    return {key: value for key, value in values.items() if key in columns}


def _product_import_header_key(value):
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(value or "").strip().lower()).split())


def _product_import_row_value(row, field):
    aliases = {_product_import_header_key(alias) for alias in PRODUCT_IMPORT_ALIASES[field]}
    for key, value in row.items():
        if _product_import_header_key(key) in aliases:
            return value
    return None


def _clean_import_text(value):
    text = str(value or "").strip()
    if not text or text.lower() in {"none", "null", "nan", "n/a", "-"}:
        return None
    return text


def _clean_import_code(value, *, upper=True, compact=False, max_len=None):
    if isinstance(value, (int, float, Decimal)) and not isinstance(value, bool):
        try:
            number = Decimal(str(value))
            if number == number.to_integral_value():
                value = str(number.quantize(Decimal("1")))
        except (InvalidOperation, ValueError):
            pass
    text = _clean_import_text(value)
    if not text:
        return None
    if compact:
        text = "".join(text.split())
    if upper:
        text = text.upper()
    if max_len:
        text = text[:max_len]
    return text


def _clean_import_package_type(value):
    text = _clean_import_text(value)
    if not text:
        return None
    normalised = normalise_package_type(text)
    return str(normalised or text).strip()[:40]


def _parse_import_decimal(value):
    text = _clean_import_text(value)
    if not text:
        return None
    try:
        return float(str(text).replace(",", ""))
    except (TypeError, ValueError):
        return None


def _parse_import_int(value):
    number = _parse_import_decimal(value)
    return int(number) if number is not None else None


def _parse_import_bool(value, default=0):
    text = _clean_import_text(value)
    if text is None:
        return default
    lowered = text.lower()
    if lowered in {"1", "true", "yes", "y", "active", "enabled", "controlled"}:
        return 1
    if lowered in {"0", "false", "no", "n", "inactive", "disabled"}:
        return 0
    return default


def _read_product_import_file(storage):
    filename = (storage.filename or "").lower()
    if filename.endswith(".csv"):
        raw = storage.read().decode("utf-8-sig")
        return [dict(row) for row in csv.DictReader(io.StringIO(raw)) if any(_clean_import_text(value) for value in row.values())]

    if filename.endswith((".xlsx", ".xlsm")):
        import openpyxl

        workbook = openpyxl.load_workbook(storage, read_only=True, data_only=True)
        try:
            sheet = workbook.active
            rows = sheet.iter_rows(values_only=True)
            header = [str(value or "").strip() for value in next(rows, [])]
            records = []
            for values in rows:
                record = {key: values[idx] if idx < len(values) else None for idx, key in enumerate(header) if key}
                if any(_clean_import_text(value) for value in record.values()):
                    records.append(record)
            return records
        finally:
            workbook.close()

    raise ValueError("Upload a .xlsx, .xlsm or .csv product file.")


def _doc_product_payload_from_import_row(row, row_number):
    customer_code = _clean_import_code(_product_import_row_value(row, "customer_code"), max_len=10) or "ALL"
    sku = _clean_import_code(_product_import_row_value(row, "sku"), max_len=100)
    product_code = _clean_import_code(_product_import_row_value(row, "product_code"), max_len=100) or sku
    barcode = _clean_import_code(_product_import_row_value(row, "barcode"), upper=False, max_len=100)
    if not sku:
        sku = product_code or barcode

    description = _clean_import_text(_product_import_row_value(row, "description"))
    if not description:
        return None, f"Row {row_number}: description is required"
    if not sku and not product_code and not barcode:
        return None, f"Row {row_number}: sku or product_code is required"

    gross_weight = _parse_import_decimal(_product_import_row_value(row, "gross_weight_kg"))
    net_weight = _parse_import_decimal(_product_import_row_value(row, "net_weight_kg"))
    payload = {
        "customer_code": customer_code,
        "sku": sku,
        "barcode": barcode,
        "product_code": product_code or sku,
        "description": description[:500],
        "commodity_code": _clean_import_code(_product_import_row_value(row, "commodity_code"), compact=True, max_len=20),
        "taric_code": _clean_import_code(_product_import_row_value(row, "taric_code"), compact=True, max_len=20),
        "cus_code": _clean_import_code(_product_import_row_value(row, "cus_code"), compact=True, max_len=20),
        "national_additional_code": _clean_import_code(_product_import_row_value(row, "national_additional_code"), compact=True, max_len=10),
        "quota_order_number": _clean_import_code(_product_import_row_value(row, "quota_order_number"), compact=True, max_len=10),
        "country_of_origin": _clean_import_code(_product_import_row_value(row, "country_of_origin"), compact=True, max_len=2),
        "country_of_preferential_origin": _clean_import_code(
            _product_import_row_value(row, "country_of_preferential_origin"),
            compact=True,
            max_len=2,
        ),
        "package_type": _clean_import_package_type(_product_import_row_value(row, "package_type")),
        "procedure_code": _clean_import_code(_product_import_row_value(row, "procedure_code"), compact=True, max_len=4),
        "additional_procedure_code": _clean_import_code(
            _product_import_row_value(row, "additional_procedure_code"),
            compact=True,
            max_len=4,
        ),
        "preference_code": _clean_import_code(_product_import_row_value(row, "preference_code"), compact=True, max_len=10),
        "ni_additional_information_codes": _clean_import_code(
            _product_import_row_value(row, "ni_additional_information_codes"),
            compact=True,
            max_len=40,
        ),
        "ni_additional_info_code": _clean_import_code(
            _product_import_row_value(row, "ni_additional_information_codes"),
            compact=True,
            max_len=40,
        ),
        "nature_of_transaction": _clean_import_code(_product_import_row_value(row, "nature_of_transaction"), compact=True, max_len=40),
        "valuation_method": _clean_import_code(_product_import_row_value(row, "valuation_method"), compact=True, max_len=5),
        "valuation_indicator": _clean_import_code(_product_import_row_value(row, "valuation_indicator"), compact=True, max_len=20),
        "gross_weight_kg": gross_weight,
        "net_weight_kg": net_weight,
        "unit_price": _parse_import_decimal(_product_import_row_value(row, "unit_price")),
        "currency": _clean_import_code(_product_import_row_value(row, "currency"), compact=True, max_len=3) or "GBP",
        "statistical_unit": _clean_import_code(_product_import_row_value(row, "statistical_unit"), compact=True, max_len=10),
        "controlled_goods": _parse_import_bool(_product_import_row_value(row, "controlled_goods"), default=0),
        "controlled_goods_type": _clean_import_code(_product_import_row_value(row, "controlled_goods_type"), compact=True, max_len=40),
        "sdi_notes": _clean_import_text(_product_import_row_value(row, "sdi_notes")),
        "requires_supplementary_unit": _parse_import_bool(_product_import_row_value(row, "requires_supplementary_unit"), default=0),
        "active": _parse_import_bool(_product_import_row_value(row, "active"), default=1),
        "weight_source": _clean_import_text(_product_import_row_value(row, "weight_source")) or ("Master Data import" if gross_weight or net_weight else None),
        "weight_sample_count": _parse_import_int(_product_import_row_value(row, "weight_sample_count")),
    }
    return payload, None


def _cursor_row_dict(cursor, row):
    if not row:
        return None
    columns = [col[0] for col in cursor.description or []]
    return dict(zip(columns, row))


def _product_import_lookup_keys(payload):
    customer_code = payload.get("customer_code") or "ALL"
    keys = []
    for field in ("sku", "product_code", "barcode"):
        value = _clean_import_text(payload.get(field))
        if value:
            keys.append((customer_code, field, value))
    return keys


def _find_doc_product_catalog_match(cursor, schema_name, columns, payload):
    select_cols = [
        col
        for col in (
            "id",
            "customer_code",
            "sku",
            "product_code",
            "barcode",
            "description",
            "commodity_code",
            "taric_code",
            "cus_code",
            "national_additional_code",
            "quota_order_number",
            "country_of_origin",
            "country_of_preferential_origin",
            "package_type",
            "procedure_code",
            "additional_procedure_code",
            "preference_code",
            "ni_additional_information_codes",
            "nature_of_transaction",
            "valuation_method",
            "valuation_indicator",
            "gross_weight_kg",
            "net_weight_kg",
            "unit_price",
            "currency",
            "statistical_unit",
            "controlled_goods_type",
            "sdi_notes",
            "weight_source",
            "weight_sample_count",
        )
        if col in columns
    ]
    if not select_cols or "customer_code" not in columns:
        return None, None

    order = " ORDER BY active DESC, id" if "active" in columns else " ORDER BY id"
    for customer_code, field, value in _product_import_lookup_keys(payload):
        if field not in columns:
            continue
        cursor.execute(
            f"""
            SELECT TOP 1 {', '.join(f'[{col}]' for col in select_cols)}
            FROM {_qualified(schema_name, 'DocProductCatalog')}
            WHERE customer_code = ? AND [{field}] = ?
            {order}
            """,
            [customer_code, value],
        )
        row = _cursor_row_dict(cursor, cursor.fetchone())
        if row:
            return row, field
    return None, None


def _values_differ(left, right, *, numeric=False):
    if left in (None, "") or right in (None, ""):
        return False
    if numeric:
        try:
            if Decimal(str(right)) <= 0:
                return False
            return round(float(left), 6) != round(float(right), 6)
        except (InvalidOperation, TypeError, ValueError):
            return str(left).strip() != str(right).strip()
    return str(left).strip().upper() != str(right).strip().upper()


def _product_import_conflict_fields(payload, existing):
    checks = [
        ("sku", "SKU", False),
        ("product_code", "product code", False),
        ("barcode", "barcode", False),
        ("description", "description", False),
        ("commodity_code", "commodity code", False),
        ("country_of_origin", "origin", False),
        ("gross_weight_kg", "gross weight", True),
        ("net_weight_kg", "net weight", True),
    ]
    conflicts = []
    for field, label, numeric in checks:
        if _values_differ(payload.get(field), existing.get(field), numeric=numeric):
            conflicts.append(label)
    return conflicts


def _product_import_value_missing(value, *, numeric=False):
    if numeric:
        try:
            return value in (None, "") or Decimal(str(value)) <= 0
        except (InvalidOperation, ValueError, TypeError):
            return True
    return _clean_import_text(value) is None


def _product_import_existing_updates(payload, existing, columns):
    assignments = []
    params = []
    updated_fields = []
    numeric_fields = {
        "gross_weight_kg",
        "net_weight_kg",
        "unit_price",
        "weight_sample_count",
    }
    for field in PRODUCT_IMPORT_BACKFILL_FIELDS:
        if field not in columns:
            continue
        value = payload.get(field)
        if value in (None, ""):
            continue
        existing_value = existing.get(field)
        numeric = field in numeric_fields
        if not _product_import_value_missing(existing_value, numeric=numeric):
            continue
        assignments.append(f"[{field}] = ?")
        params.append(value)
        updated_fields.append(field)
    return assignments, params, updated_fields


def _upsert_doc_product_catalog_import(schema_name, rows, *, update_existing=False):
    if not _table_exists(schema_name, "DocProductCatalog"):
        raise RuntimeError(f"{schema_name}.DocProductCatalog is not available yet.")

    columns = _table_columns(schema_name, "DocProductCatalog")
    required = {"customer_code", "description"}
    missing_required = sorted(required - columns)
    if missing_required:
        raise RuntimeError(f"DocProductCatalog is missing required columns: {', '.join(missing_required)}")

    conn = get_db()
    cursor = conn.cursor()
    stats = {
        "source": len(rows),
        "inserted": 0,
        "updated": 0,
        "existing": 0,
        "duplicates": 0,
        "skipped": 0,
        "warnings": [],
        "errors": [],
    }
    seen_keys = set()

    try:
        for row_number, row in enumerate(rows, start=2):
            payload, error = _doc_product_payload_from_import_row(row, row_number)
            if error:
                stats["skipped"] += 1
                stats["errors"].append(error)
                continue

            row_keys = set(_product_import_lookup_keys(payload))
            if row_keys & seen_keys:
                stats["duplicates"] += 1
                stats["skipped"] += 1
                continue

            existing, matched_field = _find_doc_product_catalog_match(cursor, schema_name, columns, payload)
            usable_payload = {key: value for key, value in payload.items() if key in columns}

            if existing:
                stats["existing"] += 1
                conflicts = _product_import_conflict_fields(payload, existing)
                if conflicts:
                    code = payload.get(matched_field) or payload.get("sku") or payload.get("product_code") or payload.get("barcode")
                    stats["warnings"].append(
                        f"Row {row_number}: existing product {code} differs in {', '.join(conflicts)}"
                    )
                    stats["skipped"] += 1
                    seen_keys.update(row_keys)
                    continue
                if update_existing:
                    assignments, params, updated_fields = _product_import_existing_updates(payload, existing, columns)
                    if assignments:
                        if "updated_at" in columns:
                            assignments.append("updated_at = SYSUTCDATETIME()")
                        params.append(existing["id"])
                        cursor.execute(
                            f"""
                            UPDATE {_qualified(schema_name, 'DocProductCatalog')}
                            SET {', '.join(assignments)}
                            WHERE id = ?
                            """,
                            params,
                        )
                        stats["updated"] += 1
                    else:
                        stats["skipped"] += 1
                else:
                    stats["skipped"] += 1
                seen_keys.update(row_keys)
                continue

            insert_cols = [col for col in usable_payload if col in columns]
            placeholders = ["?" for _ in insert_cols]
            params = [usable_payload[col] for col in insert_cols]
            if "created_at" in columns:
                insert_cols.append("created_at")
                placeholders.append("SYSUTCDATETIME()")
            cursor.execute(
                f"""
                INSERT INTO {_qualified(schema_name, 'DocProductCatalog')}
                    ({', '.join(f'[{col}]' for col in insert_cols)})
                VALUES ({', '.join(placeholders)})
                """,
                params,
            )
            stats["inserted"] += 1
            seen_keys.update(row_keys)

        conn.commit()
    except Exception:
        conn.rollback()
        raise

    return stats


def _doc_product_catalog_payload_from_form(form, columns):
    payload = {
        "customer_code": _clean_import_code(form.get("customer_code"), max_len=10) or "ALL",
        "sku": _clean_import_code(form.get("sku"), max_len=100),
        "barcode": _clean_import_code(form.get("barcode"), upper=False, max_len=100),
        "product_code": _clean_import_code(form.get("product_code"), max_len=100),
        "description": _clean_import_text(form.get("description")),
        "commodity_code": _clean_import_code(form.get("commodity_code"), compact=True, max_len=20),
        "taric_code": _clean_import_code(form.get("taric_code"), compact=True, max_len=20),
        "cus_code": _clean_import_code(form.get("cus_code"), compact=True, max_len=20),
        "national_additional_code": _clean_import_code(form.get("national_additional_code"), compact=True, max_len=10),
        "quota_order_number": _clean_import_code(form.get("quota_order_number"), compact=True, max_len=10),
        "country_of_origin": _clean_import_code(form.get("country_of_origin"), compact=True, max_len=2),
        "country_of_preferential_origin": _clean_import_code(form.get("country_of_preferential_origin"), compact=True, max_len=2),
        "package_type": _clean_import_package_type(form.get("package_type")),
        "procedure_code": _clean_import_code(form.get("procedure_code"), compact=True, max_len=4),
        "additional_procedure_code": _clean_import_code(form.get("additional_procedure_code"), compact=True, max_len=4),
        "preference_code": _clean_import_code(form.get("preference_code"), compact=True, max_len=10),
        "ni_additional_information_codes": _clean_import_code(
            form.get("ni_additional_information_codes") or form.get("ni_additional_info_code"),
            compact=True,
            max_len=40,
        ),
        "ni_additional_info_code": _clean_import_code(
            form.get("ni_additional_info_code") or form.get("ni_additional_information_codes"),
            compact=True,
            max_len=40,
        ),
        "nature_of_transaction": _clean_import_code(form.get("nature_of_transaction"), compact=True, max_len=40),
        "valuation_method": _clean_import_code(form.get("valuation_method"), compact=True, max_len=5),
        "valuation_indicator": _clean_import_code(form.get("valuation_indicator"), compact=True, max_len=20),
        "gross_weight_kg": _parse_import_decimal(form.get("gross_weight_kg")),
        "net_weight_kg": _parse_import_decimal(form.get("net_weight_kg")),
        "unit_price": _parse_import_decimal(form.get("unit_price")),
        "currency": _clean_import_code(form.get("currency"), compact=True, max_len=3) or "GBP",
        "statistical_unit": _clean_import_code(form.get("statistical_unit"), compact=True, max_len=10),
        "controlled_goods": 1 if form.get("controlled_goods") else 0,
        "controlled_goods_type": _clean_import_code(form.get("controlled_goods_type"), compact=True, max_len=40),
        "sdi_notes": _clean_import_text(form.get("sdi_notes")),
        "requires_supplementary_unit": 1 if form.get("requires_supplementary_unit") else 0,
        "active": 1 if form.get("active") else 0,
        "weight_source": _clean_import_text(form.get("weight_source")),
        "weight_sample_count": _parse_import_int(form.get("weight_sample_count")),
    }
    if not payload["description"]:
        raise ValueError("Description is required.")
    if not any(payload.get(key) for key in ("sku", "product_code", "barcode")):
        raise ValueError("SKU, product code or barcode is required.")
    return {key: value for key, value in payload.items() if key in columns}


def _selected_int_ids(raw_values):
    ids = []
    for raw_id in raw_values:
        try:
            ids.append(int(raw_id))
        except (TypeError, ValueError):
            continue
    return sorted(set(ids))


def _split_product_selection_refs(raw_values):
    selected = {"Products": [], "DocProductCatalog": []}
    for raw in raw_values:
        source, _, raw_id = str(raw or "").partition(":")
        if source not in selected:
            continue
        try:
            selected[source].append(int(raw_id))
        except (TypeError, ValueError):
            continue
    return {source: sorted(set(ids)) for source, ids in selected.items() if ids}


def _apply_table_bulk_action(cursor, schema_name, table_name, ids, action):
    if not ids:
        return 0
    if not _table_exists(schema_name, table_name):
        return 0
    columns = _table_columns(schema_name, table_name)
    affected = 0
    for start in range(0, len(ids), 500):
        chunk = ids[start:start + 500]
        placeholders = ", ".join("?" for _ in chunk)
        if action == "delete":
            cursor.execute(
                f"DELETE FROM {_qualified(schema_name, table_name)} WHERE id IN ({placeholders})",
                chunk,
            )
        else:
            active_col = _product_active_column(columns) if table_name == "Products" else ("active" if "active" in columns else None)
            if not active_col:
                continue
            assignments = [f"[{active_col}] = 0"]
            if "updated_at" in columns:
                assignments.append("updated_at = SYSUTCDATETIME()")
            cursor.execute(
                f"""
                UPDATE {_qualified(schema_name, table_name)}
                SET {', '.join(assignments)}
                WHERE id IN ({placeholders})
                """,
                chunk,
            )
        if cursor.rowcount and cursor.rowcount > 0:
            affected += cursor.rowcount
    return affected


def _apply_products_bulk_action(schema_name, raw_values, action):
    action = "delete" if action == "delete" else "deactivate"
    selected = _split_product_selection_refs(raw_values)
    conn = get_db()
    cursor = conn.cursor()
    affected = 0
    try:
        for table_name, ids in selected.items():
            affected += _apply_table_bulk_action(cursor, schema_name, table_name, ids, action)
        conn.commit()
    except Exception:
        conn.rollback()
        raise
    return affected


def _apply_partners_bulk_action(schema_name, ids, action):
    action = "delete" if action == "delete" else "deactivate"
    conn = get_db()
    cursor = conn.cursor()
    try:
        affected = _apply_table_bulk_action(cursor, schema_name, "Partners", ids, action)
        conn.commit()
    except Exception:
        conn.rollback()
        raise
    return affected


def _first_existing_column(columns, *candidates):
    lowered = {str(column).lower(): column for column in columns or set()}
    for candidate in candidates:
        found = lowered.get(candidate.lower())
        if found:
            return found
    return None


def _decimal_from_value(value):
    if value in (None, ""):
        return None
    try:
        text = str(value).strip().replace(",", "")
        if not text:
            return None
        return Decimal(text)
    except (InvalidOperation, ValueError):
        return None


def _positive_decimal(value):
    number = _decimal_from_value(value)
    return number if number is not None and number > 0 else None


def _unit_weight_decimal(total_weight, package_count):
    total = _positive_decimal(total_weight)
    packages = _positive_decimal(package_count)
    if total is None or packages is None:
        return None
    return (total / packages).quantize(Decimal("0.001"), rounding=ROUND_HALF_UP)


def _normalise_lookup_key(value):
    text = str(value or "").strip()
    if not text:
        return ""
    return " ".join(text.upper().split())


def _status_lookup_key(value):
    text = str(value or "").replace("_", " ").replace("-", " ").upper()
    return " ".join(text.split())


def _product_table_weight_columns(table_name, columns):
    if table_name == "Products":
        return (
            _first_existing_column(columns, "default_gross_weight_kg", "gross_weight_kg", "gross_mass_kg"),
            _first_existing_column(columns, "default_net_weight_kg", "net_weight_kg", "net_mass_kg"),
        )
    return (
        _first_existing_column(columns, "gross_weight_kg", "default_gross_weight_kg", "gross_mass_kg"),
        _first_existing_column(columns, "net_weight_kg", "default_net_weight_kg", "net_mass_kg"),
    )


def _historical_goods_row_is_usable(row):
    if not row:
        return False
    if _clean_import_text(row.get("goods_id")):
        return True
    if _clean_import_text(row.get("cons_dec_reference")) or _clean_import_text(row.get("dec_reference")):
        return True

    remote_ok = {
        "CREATED",
        "SUBMITTED",
        "ARRIVED",
        "AUTHORISED FOR MOVEMENT",
        "AUTHORIZED FOR MOVEMENT",
        "ACCEPTED",
        "RELEASED",
    }
    local_ok = {"CREATED", "SUBMITTED", "IMPORTED"}
    for field in ("tss_status", "cons_tss_status", "parent_tss_status"):
        if _status_lookup_key(row.get(field)) in remote_ok:
            return True
    for field in ("status", "cons_status", "parent_status"):
        if _status_lookup_key(row.get(field)) in local_ok:
            return True
    return False


def _historical_goods_weight_samples(schema_name, *, limit=PRODUCT_WEIGHT_FETCH_SAMPLE_LIMIT):
    if not _table_exists(schema_name, "StagingGoodsItems"):
        return []

    columns = _table_columns(schema_name, "StagingGoodsItems")
    key_columns = [column for column in PRODUCT_WEIGHT_LOOKUP_FIELDS if column in columns]
    if not key_columns or "number_of_packages" not in columns:
        return []
    if "gross_mass_kg" not in columns and "net_mass_kg" not in columns:
        return []

    select_columns = []
    for column in (
        "staging_id",
        *PRODUCT_WEIGHT_LOOKUP_FIELDS,
        "goods_description",
        "number_of_packages",
        "gross_mass_kg",
        "net_mass_kg",
        "goods_id",
        "status",
        "tss_status",
    ):
        if column in columns:
            select_columns.append(f"g.[{column}] AS [{column}]")

    join_sql = ""
    if "staging_cons_id" in columns and _table_exists(schema_name, "StagingConsignments"):
        cons_columns = _table_columns(schema_name, "StagingConsignments")
        if "status" in cons_columns:
            select_columns.append("c.[status] AS [cons_status]")
        if "tss_status" in cons_columns:
            select_columns.append("c.[tss_status] AS [cons_tss_status]")
        dec_col = _first_existing_column(cons_columns, "dec_reference", "consignment_reference")
        if dec_col:
            select_columns.append(f"c.[{dec_col}] AS [cons_dec_reference]")
        if any(column in cons_columns for column in ("status", "tss_status", "dec_reference", "consignment_reference")):
            join_sql = f" LEFT JOIN {_qualified(schema_name, 'StagingConsignments')} c ON c.[staging_id] = g.[staging_cons_id]"

    key_where = " OR ".join(
        f"NULLIF(LTRIM(RTRIM(CAST(g.[{column}] AS NVARCHAR(200)))), '') IS NOT NULL"
        for column in key_columns
    )
    weight_where = []
    if "gross_mass_kg" in columns:
        weight_where.append("g.[gross_mass_kg] IS NOT NULL")
    if "net_mass_kg" in columns:
        weight_where.append("g.[net_mass_kg] IS NOT NULL")
    order_col = _first_existing_column(columns, "updated_at", "submitted_at", "created_at", "staging_id")
    order_sql = f" ORDER BY g.[{order_col}] DESC" if order_col else ""
    safe_limit = max(1, min(int(limit or PRODUCT_WEIGHT_FETCH_SAMPLE_LIMIT), 50000))
    return query_all(
        f"""
        SELECT TOP {safe_limit} {', '.join(select_columns)}
        FROM {_qualified(schema_name, 'StagingGoodsItems')} g
        {join_sql}
        WHERE ({key_where})
          AND g.[number_of_packages] IS NOT NULL
          AND ({' OR '.join(weight_where)})
        {order_sql}
        """
    )


def _historical_goods_weight_index(rows):
    index = {}
    stats = {"raw_samples": len(rows or []), "usable_samples": 0}
    for row in rows or []:
        if not _historical_goods_row_is_usable(row):
            continue
        package_count = _positive_decimal(row.get("number_of_packages"))
        if package_count is None:
            continue
        gross_unit = _unit_weight_decimal(row.get("gross_mass_kg"), package_count)
        net_unit = _unit_weight_decimal(row.get("net_mass_kg"), package_count)
        if gross_unit is None and net_unit is None:
            continue

        lookup_keys = {
            _normalise_lookup_key(row.get(field))
            for field in PRODUCT_WEIGHT_LOOKUP_FIELDS
            if _normalise_lookup_key(row.get(field))
        }
        if not lookup_keys:
            continue

        stats["usable_samples"] += 1
        sample = {
            "gross_unit": gross_unit,
            "net_unit": net_unit,
            "staging_id": row.get("staging_id"),
        }
        for lookup_key in lookup_keys:
            index.setdefault(lookup_key, []).append(sample)
    return index, stats


def _median_decimal(values):
    ordered = sorted(values)
    midpoint = len(ordered) // 2
    if len(ordered) % 2:
        return ordered[midpoint]
    return ((ordered[midpoint - 1] + ordered[midpoint]) / Decimal("2")).quantize(Decimal("0.001"), rounding=ROUND_HALF_UP)


def _stable_unit_weight_summary(samples, field):
    values = [sample[field] for sample in samples or [] if sample.get(field) is not None]
    if not values:
        return {"value": None, "sample_count": 0, "conflict": False}
    median = _median_decimal(values).quantize(Decimal("0.001"), rounding=ROUND_HALF_UP)
    spread = max(values) - min(values)
    tolerance = max(Decimal("0.050"), abs(median) * Decimal("0.10"))
    return {
        "value": None if spread > tolerance else median,
        "sample_count": len(values),
        "conflict": spread > tolerance,
    }


def _load_missing_weight_product_candidates(schema_name, table_name, columns, *, limit=PRODUCT_WEIGHT_FETCH_ROW_LIMIT):
    if "id" not in columns:
        return []
    key_columns = [column for column in PRODUCT_WEIGHT_LOOKUP_FIELDS if column in columns]
    if not key_columns:
        return []

    gross_col, net_col = _product_table_weight_columns(table_name, columns)
    weight_columns = [column for column in (gross_col, net_col) if column]
    if not weight_columns:
        return []

    select_columns = []
    for column in ("id", *PRODUCT_WEIGHT_LOOKUP_FIELDS, gross_col, net_col, "weight_source", "weight_sample_count"):
        if column and column in columns and column not in select_columns:
            select_columns.append(column)

    missing_where = " OR ".join(
        f"([{column}] IS NULL OR TRY_CONVERT(decimal(18,6), [{column}]) IS NULL OR TRY_CONVERT(decimal(18,6), [{column}]) <= 0)"
        for column in weight_columns
    )
    key_where = " OR ".join(
        f"NULLIF(LTRIM(RTRIM(CAST([{column}] AS NVARCHAR(200)))), '') IS NOT NULL"
        for column in key_columns
    )
    where_parts = [f"({missing_where})", f"({key_where})"]
    active_col = _product_active_column(columns) if table_name == "Products" else ("active" if "active" in columns else None)
    if active_col:
        where_parts.append(f"COALESCE([{active_col}], 1) = 1")

    order_col = _first_existing_column(columns, "updated_at", "created_at", "id")
    order_sql = f" ORDER BY [{order_col}] DESC" if order_col else ""
    safe_limit = max(1, min(int(limit or PRODUCT_WEIGHT_FETCH_ROW_LIMIT), 5000))
    return query_all(
        f"""
        SELECT TOP {safe_limit} {', '.join(f'[{column}]' for column in select_columns)}
        FROM {_qualified(schema_name, table_name)}
        WHERE {' AND '.join(where_parts)}
        {order_sql}
        """
    )


def _history_samples_for_product(row, history_index):
    for field in PRODUCT_WEIGHT_LOOKUP_FIELDS:
        lookup_key = _normalise_lookup_key((row or {}).get(field))
        if lookup_key and lookup_key in history_index:
            return lookup_key, history_index[lookup_key]
    return "", []


def _product_weight_fetch_updates_from_candidates(table_name, columns, candidates, history_index):
    gross_col, net_col = _product_table_weight_columns(table_name, columns)
    source_col = _first_existing_column(columns, "weight_source")
    sample_count_col = _first_existing_column(columns, "weight_sample_count")
    updates = []
    stats = {
        "checked": len(candidates or []),
        "matched": 0,
        "updates": 0,
        "no_match": 0,
        "conflicts": 0,
        "no_safe_weight": 0,
    }
    for row in candidates or []:
        lookup_key, samples = _history_samples_for_product(row, history_index)
        if not samples:
            stats["no_match"] += 1
            continue

        stats["matched"] += 1
        assignments = []
        sample_count = 0
        if gross_col and _positive_decimal(row.get(gross_col)) is None:
            summary = _stable_unit_weight_summary(samples, "gross_unit")
            if summary["conflict"]:
                stats["conflicts"] += 1
            elif summary["value"] is not None:
                assignments.append((gross_col, summary["value"]))
                sample_count = max(sample_count, summary["sample_count"])
        if net_col and _positive_decimal(row.get(net_col)) is None:
            summary = _stable_unit_weight_summary(samples, "net_unit")
            if summary["conflict"]:
                stats["conflicts"] += 1
            elif summary["value"] is not None:
                assignments.append((net_col, summary["value"]))
                sample_count = max(sample_count, summary["sample_count"])

        if not assignments:
            stats["no_safe_weight"] += 1
            continue

        if source_col:
            assignments.append((source_col, PRODUCT_WEIGHT_FETCH_SOURCE))
        if sample_count_col:
            assignments.append((sample_count_col, sample_count))
        updates.append({
            "table": table_name,
            "id": row.get("id"),
            "lookup_key": lookup_key,
            "assignments": assignments,
            "touch_updated_at": "updated_at" in columns,
        })
        stats["updates"] += 1
    return updates, stats


def _build_product_weight_fetch_updates(schema_name):
    history_rows = _historical_goods_weight_samples(schema_name)
    history_index, history_stats = _historical_goods_weight_index(history_rows)
    updates = []
    stats = {
        **history_stats,
        "checked": 0,
        "matched": 0,
        "updates": 0,
        "no_match": 0,
        "conflicts": 0,
        "no_safe_weight": 0,
        "tables": {},
    }
    if not history_index:
        return updates, stats

    for table_name in ("Products", "DocProductCatalog"):
        if not _table_exists(schema_name, table_name):
            continue
        columns = _table_columns(schema_name, table_name)
        candidates = _load_missing_weight_product_candidates(schema_name, table_name, columns)
        table_updates, table_stats = _product_weight_fetch_updates_from_candidates(
            table_name,
            columns,
            candidates,
            history_index,
        )
        stats["tables"][table_name] = table_stats
        for key in ("checked", "matched", "updates", "no_match", "conflicts", "no_safe_weight"):
            stats[key] += table_stats[key]
        updates.extend(table_updates)
    return updates, stats


def _apply_product_weight_fetch_updates(schema_name, updates):
    if not updates:
        return 0
    conn = get_db()
    cursor = conn.cursor()
    affected = 0
    try:
        for update in updates:
            assignments = update["assignments"]
            set_sql = ", ".join(f"[{column}] = ?" for column, _value in assignments)
            if update.get("touch_updated_at"):
                set_sql += ", [updated_at] = SYSUTCDATETIME()"
            params = [value for _column, value in assignments] + [update["id"]]
            cursor.execute(
                f"UPDATE {_qualified(schema_name, update['table'])} SET {set_sql} WHERE [id] = ?",
                params,
            )
            if cursor.rowcount and cursor.rowcount > 0:
                affected += cursor.rowcount
        conn.commit()
    except Exception:
        conn.rollback()
        raise
    return affected


def _flash_products_table_missing(schema_name):
    flash(
        f"{schema_name}.Products is not present yet. Run the tenant master data migration to enable product master data.",
        "warning",
    )


def _flash_partners_table_missing(schema_name):
    flash(
        f"{schema_name}.Partners is not present yet. Run the tenant master data migration to enable tenant partner master data.",
        "warning",
    )


def _cv_table_names():
    try:
        rows = query_all(
            """
            SELECT TABLE_NAME
            FROM INFORMATION_SCHEMA.TABLES
            WHERE TABLE_SCHEMA = 'TSS' AND TABLE_NAME LIKE 'CV[_]%'
            ORDER BY TABLE_NAME
            """
        )
        names = [r["TABLE_NAME"] for r in rows]
        return names or STATIC_CV_TABLE_NAMES
    except Exception:
        return STATIC_CV_TABLE_NAMES


def _load_cv_options(table_name, value_col="value", label_col="name"):
    if table_name not in set(_cv_table_names()) | set(STATIC_CV_TABLE_NAMES):
        return []
    try:
        return query_all(
            f"SELECT [{value_col}] AS value, [{label_col}] AS name FROM TSS.[{table_name}] ORDER BY [{label_col}]"
        )
    except Exception:
        return []


def _product_form_choices():
    return {
        "countries": _load_cv_options("CV_country"),
        "package_types": _load_cv_options("CV_type_of_package"),
        "procedure_codes": _load_cv_options("CV_procedure_code"),
        "valuation_methods": _load_cv_options("CV_valuation_method"),
        "preferences": _load_cv_options("CV_preference"),
        "ni_codes": _load_cv_options("CV_ni_additional_information_code"),
    }


def _to_float(value):
    value = (value or "").strip()
    if not value:
        return None
    try:
        return float(value)
    except ValueError:
        return None


def _load_config_rows(schema_name, category):
    if not _table_exists(schema_name, "AppConfiguration"):
        return []
    try:
        return query_all(
            f"""
            SELECT category, config_key, config_value
            FROM {_qualified(schema_name, 'AppConfiguration')}
            WHERE category = ?
            ORDER BY config_key
            """,
            [category],
        )
    except Exception:
        return []


def _load_config_map(schema_name, category):
    return {row["config_key"]: row.get("config_value") or "" for row in _load_config_rows(schema_name, category)}


def _resolve_tenant_config_value(raw_map, category, key, fallback=""):
    raw_value = (raw_map.get(key) or "").strip()
    if raw_value:
        return raw_value, "tenant"

    env_value = (os.environ.get(f"{category}_{key}", "") or "").strip()
    if env_value:
        return env_value, "env"

    if fallback not in (None, ""):
        return str(fallback), "default"
    return "", "empty"


def _build_ingest_default_sections(schema_name):
    raw_map = _load_config_map(schema_name, "INGEST_AUTO")
    sections = []
    for section in INGEST_DEFAULT_SECTIONS:
        built_fields = []
        for field in section["fields"]:
            effective_value, source = _resolve_tenant_config_value(
                raw_map,
                "INGEST_AUTO",
                field["key"],
                fallback=field.get("fallback", ""),
            )
            built_fields.append({
                **field,
                "input_name": f"cfg_INGEST_AUTO__{field['key']}",
                "raw_value": raw_map.get(field["key"], ""),
                "effective_value": effective_value,
                "source": source,
            })
        sections.append({**section, "fields": built_fields})
    return sections


def _config_source_label(source):
    return {
        "tenant": "Tenant override",
        "env": "ENV fallback",
        "default": "App default",
        "empty": "Blank",
    }.get(source, "Resolved")


def _truthy(value):
    return str(value or "").strip().lower() in {"1", "true", "yes", "on", "enabled"}


def _build_validation_settings(schema_name):
    raw_map = _load_config_map(schema_name, "VALIDATION")
    settings = []
    for field in VALIDATION_SETTINGS:
        effective_value, source = _resolve_tenant_config_value(
            raw_map,
            field["category"],
            field["key"],
            fallback=field.get("fallback", ""),
        )
        settings.append({
            **field,
            "input_name": f"cfg_{field['category']}__{field['key']}",
            "raw_value": raw_map.get(field["key"], ""),
            "effective_value": effective_value,
            "enabled": _truthy(effective_value),
            "source": source,
        })
    return settings


def _flash_message(category, text, technical_url=None, technical_label="Technical"):
    payload = {"text": text}
    if technical_url:
        payload["technical_url"] = technical_url
        payload["technical_label"] = technical_label
    flash(payload, category)


def _log_master_data_event(schema_name, call_type, response_message, *, error_detail=None, request_payload=None):
    return insert_api_call_log(
        schema_name,
        call_type,
        http_method=request.method,
        url=request.full_path if request.query_string else request.path,
        request_payload=request_payload,
        http_status=0,
        response_status="error",
        response_message=response_message,
        error_detail=error_detail or response_message,
    )


def _technical_api_url(log_id):
    return url_for("technical.index", tab="api", log_id=log_id, _anchor=f"api-log-{log_id}")


def _save_config_updates(schema_name, tenant_code, updates):
    if not updates:
        return
    if not _table_exists(schema_name, "AppConfiguration"):
        flash(
            f"{schema_name}.AppConfiguration is not available yet, so tenant operational defaults could not be saved.",
            "warning",
        )
        return

    field_meta = _tenant_config_field_map()
    conn = get_db()
    cursor = conn.cursor()
    errors = []

    for category, items in updates.items():
        for config_key, value in items.items():
            text = (value or "").strip()
            description = field_meta.get(config_key, {}).get("help", "")
            try:
                cursor.execute(
                    f"""
                    UPDATE {_qualified(schema_name, 'AppConfiguration')}
                    SET config_value = ?, updated_at = GETUTCDATE()
                    WHERE category = ? AND config_key = ?
                    """,
                    text,
                    category,
                    config_key,
                )
                if cursor.rowcount == 0:
                    cursor.execute(
                        f"""
                        INSERT INTO {_qualified(schema_name, 'AppConfiguration')}
                            (category, config_key, config_value, description, is_secret)
                        VALUES (?, ?, ?, ?, 0)
                        """,
                        category,
                        config_key,
                        text,
                        description,
                    )
            except Exception as exc:
                errors.append(f"{category}.{config_key}: {exc}")

    if errors:
        conn.rollback()
        raise RuntimeError("; ".join(errors))

    conn.commit()
    config_store.reload(tenant_code)


def _extract_config_updates(form):
    updates = {}
    for key, value in form.items():
        if not key.startswith("cfg_"):
            continue
        try:
            category, config_key = key[4:].split("__", 1)
        except ValueError:
            continue
        updates.setdefault(category, {})[config_key] = value
    return updates


def _upsert_company(form, schema_name=None):
    schema_name = schema_name or _tenant_ctx()[1]
    company = _company_record(schema_name)
    values = [
        form.get("company_name"),
        form.get("trading_name"),
        form.get("eori_xi"),
        form.get("eori_gb"),
        form.get("address_line1"),
        form.get("address_line2"),
        form.get("city"),
        form.get("postcode"),
        form.get("country"),
        form.get("contact_name"),
        form.get("contact_email"),
        form.get("contact_phone"),
        form.get("ukims_authorisation"),
        form.get("scdp_authorisation"),
    ]

    if company:
        execute(
            """
            UPDATE {table_name} SET
                company_name=?,
                trading_name=?,
                eori_xi=?,
                eori_gb=?,
                address_line1=?,
                address_line2=?,
                city=?,
                postcode=?,
                country=?,
                contact_name=?,
                contact_email=?,
                contact_phone=?,
                ukims_authorisation=?,
                scdp_authorisation=?,
                updated_at=GETUTCDATE()
            WHERE id = ?
            """.format(table_name=_qualified(schema_name, "CompanyMaster")),
            values + [company["id"]],
        )
    else:
        execute(
            """
            INSERT INTO {table_name} (
                company_name, trading_name, eori_xi, eori_gb,
                address_line1, address_line2, city, postcode, country,
                contact_name, contact_email, contact_phone,
                ukims_authorisation, scdp_authorisation
            ) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?)
            """.format(table_name=_qualified(schema_name, "CompanyMaster")),
            values,
        )


@master_data_bp.route("/")
def index():
    tenant, schema_name = _tenant_ctx()
    company = _company_record(schema_name)
    partner_count = _safe_count(f"SELECT COUNT(*) AS cnt FROM {_qualified(schema_name, 'Partners')} WHERE active = 1")
    has_products = _table_exists(schema_name, "Products")
    product_columns = _table_columns(schema_name, "Products") if has_products else set()
    product_active_col = _product_active_column(product_columns)
    product_where = f" WHERE {product_active_col} = 1" if product_active_col else ""
    product_count = (
        _safe_count(f"SELECT COUNT(*) AS cnt FROM {_qualified(schema_name, 'Products')}{product_where}")
        if has_products
        else None
    )
    catalog_count = _doc_product_catalog_count(schema_name)
    visible_product_count = product_count if product_count else catalog_count
    cv_tables = _cv_table_names()
    return render_template(
        "master_data/index.html",
        company=company,
        partner_count=partner_count or 0,
        product_count=visible_product_count,
        products_table_count=product_count,
        catalog_count=catalog_count,
        data_counts=_raw_data_counts(),
        eori_cache_stats=_eori_cache_stats(schema_name),
        cv_table_count=len(cv_tables),
        tenant=tenant,
        validation_settings=_build_validation_settings(schema_name),
        has_app_config=_table_exists(schema_name, "AppConfiguration"),
        config_source_label=_config_source_label,
    )


@master_data_bp.route("/validation-settings", methods=["POST"])
def validation_settings_save():
    tenant, schema_name = _tenant_ctx()
    strict_enabled = "true" if request.form.get("STRICT_MASTERDATA_VALIDATION") == "true" else "false"
    try:
        _save_config_updates(
            schema_name,
            tenant["code"],
            {"VALIDATION": {"STRICT_MASTERDATA_VALIDATION": strict_enabled}},
        )
        state = "enabled" if strict_enabled == "true" else "disabled"
        flash(f"Strict local masterdata validation {state}.", "success")
    except Exception as exc:
        message = f"Error saving validation settings: {exc}"
        log_id = _log_master_data_event(
            schema_name,
            "MASTERDATA_VALIDATION_SETTINGS_SAVE",
            message,
            error_detail=exc,
            request_payload=dict(request.form),
        )
        _flash_message("danger", message, technical_url=_technical_api_url(log_id) if log_id else None)
    return redirect(url_for("master_data.index"))


@master_data_bp.route("/company")
def company():
    tenant, schema_name = _tenant_ctx()
    return render_template(
        "master_data/company.html",
        company=_company_record(schema_name),
        tenant=tenant,
    )


@master_data_bp.route("/company/edit", methods=["GET", "POST"])
def company_edit():
    tenant, schema_name = _tenant_ctx()
    company = _company_record(schema_name) or {}
    advanced_sections = _build_ingest_default_sections(schema_name)
    advanced_requested = str(request.args.get("advanced") or "").strip().lower() in {"1", "true", "yes"}
    highlight_config_key = (request.args.get("highlight") or "").strip().upper()
    if request.method == "POST":
        try:
            _upsert_company(request.form, schema_name)
            _save_config_updates(schema_name, tenant["code"], _extract_config_updates(request.form))
            flash("Company details saved.", "success")
            return redirect(url_for("master_data.company"))
        except Exception as exc:
            message = f"Error saving company details: {exc}"
            log_id = _log_master_data_event(
                schema_name,
                "MASTERDATA_COMPANY_SAVE",
                message,
                error_detail=exc,
                request_payload=dict(request.form),
            )
            _flash_message("danger", message, technical_url=_technical_api_url(log_id) if log_id else None)
            company = dict(company)
            company.update(request.form)
            advanced_sections = _build_ingest_default_sections(schema_name)
            for section in advanced_sections:
                for field in section["fields"]:
                    field["raw_value"] = request.form.get(field["input_name"], field["raw_value"])
                    field["effective_value"], field["source"] = _resolve_tenant_config_value(
                        {field["key"]: field["raw_value"]},
                        "INGEST_AUTO",
                        field["key"],
                        fallback=field.get("fallback", ""),
                    )
    return render_template(
        "master_data/company_edit.html",
        company=company,
        tenant=tenant,
        advanced_sections=advanced_sections,
        advanced_open=request.method == "POST" or advanced_requested,
        highlight_config_key=highlight_config_key,
        has_app_config=_table_exists(schema_name, "AppConfiguration"),
        config_source_label=_config_source_label,
    )


@master_data_bp.route("/company/save", methods=["POST"])
def company_save():
    tenant, schema_name = _tenant_ctx()
    try:
        _upsert_company(request.form, schema_name)
        _save_config_updates(schema_name, tenant["code"], _extract_config_updates(request.form))
        flash("Company details saved.", "success")
    except Exception as exc:
        message = f"Error saving company details: {exc}"
        log_id = _log_master_data_event(
            schema_name,
            "MASTERDATA_COMPANY_SAVE",
            message,
            error_detail=exc,
            request_payload=dict(request.form),
        )
        _flash_message("danger", message, technical_url=_technical_api_url(log_id) if log_id else None)
    return redirect(url_for("master_data.company"))


@master_data_bp.route("/partners")
def partners():
    tenant, schema_name = _tenant_ctx()
    search = (request.args.get("q") or "").strip()
    partner_type = (request.args.get("type") or "").strip()
    where = []
    params = []
    if not _table_exists(schema_name, "Partners"):
        _flash_partners_table_missing(schema_name)
        return render_template(
            "master_data/partners.html",
            partners=[],
            search=search,
            partner_type=partner_type,
            partner_types=PARTNER_TYPES,
            tenant=tenant,
        )

    if search:
        where.append("(partner_name LIKE ? OR eori LIKE ? OR city LIKE ? OR postcode LIKE ?)")
        like = f"%{search}%"
        params.extend([like, like, like, like])
    if partner_type in PARTNER_TYPES:
        where.append("partner_type = ?")
        params.append(partner_type)

    sql = f"SELECT * FROM {_qualified(schema_name, 'Partners')}"
    if where:
        sql += " WHERE " + " AND ".join(where)
    sql += " ORDER BY active DESC, partner_type, partner_name"

    partners_list = query_all(sql, params)
    return render_template(
        "master_data/partners.html",
        partners=partners_list,
        search=search,
        partner_type=partner_type,
        partner_types=PARTNER_TYPES,
        tenant=tenant,
    )


@master_data_bp.route("/partners/bulk-action", methods=["POST"])
def partners_bulk_action():
    tenant, schema_name = _tenant_ctx()
    redirect_args = {}
    search = (request.form.get("q") or "").strip()
    partner_type = (request.form.get("type") or "").strip()
    if search:
        redirect_args["q"] = search
    if partner_type:
        redirect_args["type"] = partner_type

    ids = _selected_int_ids(request.form.getlist("selected_ids"))
    action = "delete" if request.form.get("bulk_action") == "delete" else "deactivate"
    if not ids:
        flash("Select at least one partner first.", "warning")
        return redirect(url_for("master_data.partners", **redirect_args))

    try:
        affected = _apply_partners_bulk_action(schema_name, ids, action)
        verb = "Deleted" if action == "delete" else "Deactivated"
        flash(f"{verb} {affected} selected partner{'s' if affected != 1 else ''}.", "success" if affected else "warning")
    except Exception as exc:
        message = f"Error applying partner bulk action: {exc}"
        log_id = _log_master_data_event(
            schema_name,
            "MASTERDATA_PARTNERS_BULK_ACTION",
            message,
            error_detail=exc,
            request_payload={"ids": ids, "action": action, "tenant": tenant["code"]},
        )
        _flash_message("danger", message, technical_url=_technical_api_url(log_id) if log_id else None)
    return redirect(url_for("master_data.partners", **redirect_args))


@master_data_bp.route("/partners/new", methods=["GET", "POST"])
@master_data_bp.route("/partners/create", methods=["GET", "POST"])
def partner_create():
    tenant, schema_name = _tenant_ctx()
    if not _table_exists(schema_name, "Partners"):
        _flash_partners_table_missing(schema_name)
        return redirect(url_for("master_data.partners"))
    partner = {}
    if request.method == "POST":
        partner = dict(request.form)
        try:
            execute(
                """
                INSERT INTO {table_name} (
                    partner_type, partner_name, eori, address_line1, city, postcode,
                    country, contact_name, contact_email, contact_phone, active,
                    created_at, updated_at
                ) VALUES (?,?,?,?,?,?,?,?,?,?,?,GETUTCDATE(),GETUTCDATE())
                """.format(table_name=_qualified(schema_name, "Partners")),
                [
                    request.form.get("partner_type"),
                    request.form.get("partner_name"),
                    request.form.get("eori") or None,
                    request.form.get("address_line1") or None,
                    request.form.get("city") or None,
                    request.form.get("postcode") or None,
                    request.form.get("country") or None,
                    request.form.get("contact_name") or None,
                    request.form.get("contact_email") or None,
                    request.form.get("contact_phone") or None,
                    1 if request.form.get("active") else 0,
                ],
            )
            flash(f'Partner "{request.form.get("partner_name")}" created.', "success")
            return redirect(url_for("master_data.partners"))
        except Exception as exc:
            message = f"Error creating partner: {exc}"
            log_id = _log_master_data_event(
                schema_name,
                "MASTERDATA_PARTNER_CREATE",
                message,
                error_detail=exc,
                request_payload=dict(request.form),
            )
            _flash_message("danger", message, technical_url=_technical_api_url(log_id) if log_id else None)
    return render_template(
        "master_data/partner_form.html",
        partner=partner,
        mode="create",
        partner_types=PARTNER_TYPES,
        tenant=tenant,
    )


@master_data_bp.route("/partners/<int:partner_id>/edit", methods=["GET", "POST"])
def partner_edit(partner_id):
    tenant, schema_name = _tenant_ctx()
    if not _table_exists(schema_name, "Partners"):
        _flash_partners_table_missing(schema_name)
        return redirect(url_for("master_data.partners"))
    partner = query_one(f"SELECT * FROM {_qualified(schema_name, 'Partners')} WHERE id = ?", [partner_id]) or {}
    if not partner:
        flash("Partner not found.", "warning")
        return redirect(url_for("master_data.partners"))

    if request.method == "POST":
        partner = dict(partner)
        partner.update(request.form)
        try:
            execute(
                """
                UPDATE {table_name} SET
                    partner_type=?,
                    partner_name=?,
                    eori=?,
                    address_line1=?,
                    city=?,
                    postcode=?,
                    country=?,
                    contact_name=?,
                    contact_email=?,
                    contact_phone=?,
                    active=?,
                    updated_at=GETUTCDATE()
                WHERE id = ?
                """.format(table_name=_qualified(schema_name, "Partners")),
                [
                    request.form.get("partner_type"),
                    request.form.get("partner_name"),
                    request.form.get("eori") or None,
                    request.form.get("address_line1") or None,
                    request.form.get("city") or None,
                    request.form.get("postcode") or None,
                    request.form.get("country") or None,
                    request.form.get("contact_name") or None,
                    request.form.get("contact_email") or None,
                    request.form.get("contact_phone") or None,
                    1 if request.form.get("active") else 0,
                    partner_id,
                ],
            )
            flash("Partner updated.", "success")
            return redirect(url_for("master_data.partners"))
        except Exception as exc:
            message = f"Error updating partner: {exc}"
            log_id = _log_master_data_event(
                schema_name,
                "MASTERDATA_PARTNER_UPDATE",
                message,
                error_detail=exc,
                request_payload=dict(request.form),
            )
            _flash_message("danger", message, technical_url=_technical_api_url(log_id) if log_id else None)
    return render_template(
        "master_data/partner_form.html",
        partner=partner,
        mode="edit",
        partner_types=PARTNER_TYPES,
        tenant=tenant,
    )


@master_data_bp.route("/products")
def products():
    tenant, active_schema_name = _tenant_ctx()
    search = (request.args.get("q") or "").strip()
    info_status = _normalize_product_info_status(request.args.get("info_status"))
    page_size = 100
    product_tenant_options = _product_tenant_options(tenant)
    selected_product_tenant, product_scope_tenants, product_scope_label = _product_scope(
        tenant,
        request.args.get("product_tenant"),
    )
    all_tenant_products = selected_product_tenant == "ALL"
    show_all_allowed = not all_tenant_products
    show_all = request.args.get("show_all") == "1" and show_all_allowed
    if all_tenant_products:
        schema_name = active_schema_name
    else:
        schema_name = product_scope_tenants[0]["schema"]
    products_read_only = bool(product_tenant_options and selected_product_tenant == "ALL")
    try:
        page = max(1, int(request.args.get("page") or 1))
    except (TypeError, ValueError):
        page = 1
    try:
        if all_tenant_products:
            loaded = _load_product_master_rows_for_tenants(
                product_scope_tenants,
                search,
                info_status=info_status,
            )
        else:
            loaded = _load_product_master_rows(schema_name, search, info_status=info_status)
            if product_tenant_options:
                loaded = dict(loaded)
                loaded["rows"] = _annotate_product_rows_for_tenant(loaded["rows"], product_scope_tenants[0])
    except Exception as exc:
        loaded = {
            "rows": [],
            "has_products_table": _table_exists(schema_name, "Products"),
            "has_catalog_table": _table_exists(schema_name, "DocProductCatalog"),
            "has_product_source": False,
            "product_source_label": "Products",
        }
        message = f"Error loading products: {exc}"
        log_id = _log_master_data_event(
            schema_name,
            "MASTERDATA_PRODUCTS_LIST",
            message,
            error_detail=exc,
            request_payload={
                "q": search,
                "info_status": info_status,
                "product_tenant": selected_product_tenant,
            },
        )
        _flash_message("danger", message, technical_url=_technical_api_url(log_id) if log_id else None)
    if not loaded["has_product_source"]:
        if all_tenant_products:
            flash("No product master tables were found across the selected tenants.", "warning")
        else:
            _flash_products_table_missing(schema_name)

    all_products = loaded["rows"]
    total_products = len(all_products)
    total_pages = max(1, math.ceil(total_products / page_size)) if total_products else 1
    if page > total_pages:
        page = total_pages
    if show_all:
        visible_products = all_products
    else:
        start = (page - 1) * page_size
        visible_products = all_products[start:start + page_size]

    return render_template(
        "master_data/products.html",
        products=visible_products,
        total_products=total_products,
        page=page,
        page_size=page_size,
        total_pages=total_pages,
        show_all=show_all,
        search=search,
        info_status=info_status,
        product_info_status_options=PRODUCT_INFO_STATUS_OPTIONS,
        product_tenant_options=product_tenant_options,
        selected_product_tenant=selected_product_tenant,
        product_scope_label=product_scope_label,
        show_tenant_column=bool(product_tenant_options),
        products_read_only=products_read_only,
        show_all_allowed=show_all_allowed,
        has_products_table=loaded["has_products_table"],
        has_catalog_table=loaded["has_catalog_table"],
        has_product_source=loaded["has_product_source"],
        product_source_label=loaded["product_source_label"],
        catalog_count=(
            _doc_product_catalog_count_for_tenants(product_scope_tenants)
            if all_tenant_products
            else _doc_product_catalog_count(schema_name)
        ),
        tenant=tenant,
    )


@master_data_bp.route("/products/export.csv")
def products_export_csv():
    tenant, active_schema_name = _tenant_ctx()
    product_tenant_options = _product_tenant_options(tenant)
    selected_product_tenant, product_scope_tenants, _scope_label = _product_scope(
        tenant,
        request.args.get("product_tenant"),
    )
    all_tenant_products = selected_product_tenant == "ALL"
    if all_tenant_products:
        flash("Choose one tenant before exporting product masterdata.", "warning")
        return redirect(url_for("master_data.products", product_tenant="ALL"))
    schema_name = active_schema_name if all_tenant_products else product_scope_tenants[0]["schema"]
    try:
        loaded = _load_product_master_rows(schema_name, "", catalog_active_only=False)
        rows = (
            _annotate_product_rows_for_tenant(loaded["rows"], product_scope_tenants[0])
            if product_tenant_options
            else loaded["rows"]
        )
    except Exception as exc:
        message = f"Error exporting products: {exc}"
        log_id = _log_master_data_event(
            schema_name,
            "MASTERDATA_PRODUCTS_EXPORT",
            message,
            error_detail=exc,
        )
        _flash_message("danger", message, technical_url=_technical_api_url(log_id) if log_id else None)
        return redirect(url_for("master_data.products", product_tenant=selected_product_tenant) if product_tenant_options else url_for("master_data.products"))

    output = io.StringIO()
    export_fields = (
        ["tenant_code", "tenant_name"] + PRODUCT_EXPORT_FIELDS
        if product_tenant_options
        else PRODUCT_EXPORT_FIELDS
    )
    writer = csv.DictWriter(output, fieldnames=export_fields, extrasaction="ignore")
    writer.writeheader()
    for row in rows:
        writer.writerow({field: row.get(field) for field in export_fields})

    filename_prefix = selected_product_tenant.lower() if product_tenant_options else tenant["code"].lower()
    filename = f"{filename_prefix}_products_masterdata.csv"
    return Response(
        output.getvalue(),
        mimetype="text/csv",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@master_data_bp.route("/products/import", methods=["POST"])
def products_import():
    tenant, schema_name = _product_target_ctx()
    upload = request.files.get("product_file")
    if not upload or not upload.filename:
        flash("Choose an Excel or CSV file to import.", "warning")
        return redirect(_products_url_for_tenant(tenant))

    try:
        rows = _read_product_import_file(upload)
        stats = _upsert_doc_product_catalog_import(
            schema_name,
            rows,
            update_existing=bool(request.form.get("update_existing")),
        )
        already_existing = stats["existing"] + stats["duplicates"]
        warning_count = len(stats["warnings"]) + len(stats["errors"])
        details = (
            f"{stats['inserted']} new products added, "
            f"{stats['updated']} existing products updated, "
            f"{already_existing} already existed, "
            f"{warning_count} warning{'s' if warning_count != 1 else ''}"
        )
        first_warning = (stats["warnings"] or stats["errors"] or [None])[0]
        if first_warning:
            details += f". First warning: {first_warning}"
        flash(f"Product import complete: {details}.", "success" if stats["inserted"] else "warning")
    except Exception as exc:
        message = f"Error importing products: {exc}"
        log_id = _log_master_data_event(
            schema_name,
            "MASTERDATA_PRODUCTS_IMPORT",
            message,
            error_detail=exc,
            request_payload={"filename": upload.filename, "tenant": tenant["code"]},
        )
        _flash_message("danger", message, technical_url=_technical_api_url(log_id) if log_id else None)
        return redirect(_products_url_for_tenant(tenant))

    return redirect(_products_url_for_tenant(tenant))


@master_data_bp.route("/products/fetch-info", methods=["POST"])
def products_fetch_info():
    active_tenant = get_tenant()
    selected_product_tenant = _selected_product_tenant(
        active_tenant,
        request.form.get("product_tenant") or active_tenant.get("code"),
    )
    redirect_args = {}
    if active_tenant.get("code") == "SYD" or request.form.get("product_tenant"):
        redirect_args["product_tenant"] = selected_product_tenant
    for key in ("q", "info_status", "page"):
        value = (request.form.get(key) or "").strip()
        if value:
            redirect_args[key] = value

    if selected_product_tenant == "ALL":
        flash("Choose one tenant before fetching product weights from historical goods.", "warning")
        return redirect(url_for("master_data.products", **redirect_args))

    tenant, schema_name = _product_target_ctx()
    preview_only = request.form.get("fetch_mode") != "apply"
    try:
        updates, stats = _build_product_weight_fetch_updates(schema_name)
        if preview_only:
            category = "success" if updates else "warning"
            flash(
                "Product fetch preview: "
                f"{len(updates)} product row{'s' if len(updates) != 1 else ''} can receive missing unit weights "
                f"from {stats['usable_samples']} usable historical goods sample{'s' if stats['usable_samples'] != 1 else ''}. "
                f"{stats['conflicts']} conflicting weight field{'s' if stats['conflicts'] != 1 else ''} skipped.",
                category,
            )
        else:
            affected = _apply_product_weight_fetch_updates(schema_name, updates)
            category = "success" if affected else "warning"
            flash(
                "Product fetch complete: "
                f"updated {affected} product row{'s' if affected != 1 else ''} "
                f"from {stats['usable_samples']} usable historical goods sample{'s' if stats['usable_samples'] != 1 else ''}. "
                f"{stats['conflicts']} conflicting weight field{'s' if stats['conflicts'] != 1 else ''} skipped.",
                category,
            )
    except Exception as exc:
        message = f"Error fetching product info from historical goods: {exc}"
        log_id = _log_master_data_event(
            schema_name,
            "MASTERDATA_PRODUCTS_FETCH_INFO",
            message,
            error_detail=exc,
            request_payload={"tenant": tenant["code"], "preview_only": preview_only},
        )
        _flash_message("danger", message, technical_url=_technical_api_url(log_id) if log_id else None)

    return redirect(url_for("master_data.products", **redirect_args))


@master_data_bp.route("/products/bulk-action", methods=["POST"])
def products_bulk_action():
    tenant, schema_name = _product_target_ctx()
    redirect_args = {}
    search = (request.form.get("q") or "").strip()
    info_status = _normalize_product_info_status(request.form.get("info_status"))
    show_all = request.form.get("show_all") == "1"
    selected_product_tenant = _selected_product_tenant(tenant, request.form.get("product_tenant") or tenant["code"])
    if request.form.get("product_tenant") or tenant["code"] == "SYD":
        redirect_args["product_tenant"] = selected_product_tenant
    try:
        page = max(1, int(request.form.get("page") or 1))
    except (TypeError, ValueError):
        page = 1
    if search:
        redirect_args["q"] = search
    if info_status:
        redirect_args["info_status"] = info_status
    if show_all:
        redirect_args["show_all"] = 1
    elif page > 1:
        redirect_args["page"] = page

    raw_values = request.form.getlist("selected_ids")
    selected = _split_product_selection_refs(raw_values)
    action = "delete" if request.form.get("bulk_action") == "delete" else "deactivate"
    if selected_product_tenant == "ALL":
        flash("Choose one tenant before applying bulk product changes.", "warning")
        return redirect(url_for("master_data.products", **redirect_args))
    if not selected:
        flash("Select at least one product first.", "warning")
        return redirect(url_for("master_data.products", **redirect_args))

    try:
        affected = _apply_products_bulk_action(schema_name, raw_values, action)
        verb = "Deleted" if action == "delete" else "Deactivated"
        flash(f"{verb} {affected} selected product{'s' if affected != 1 else ''}.", "success" if affected else "warning")
    except Exception as exc:
        message = f"Error applying product bulk action: {exc}"
        log_id = _log_master_data_event(
            schema_name,
            "MASTERDATA_PRODUCTS_BULK_ACTION",
            message,
            error_detail=exc,
            request_payload={"selected": raw_values, "action": action, "tenant": tenant["code"]},
        )
        _flash_message("danger", message, technical_url=_technical_api_url(log_id) if log_id else None)
    return redirect(url_for("master_data.products", **redirect_args))


@master_data_bp.route("/products/new", methods=["GET", "POST"])
@master_data_bp.route("/products/create", methods=["GET", "POST"])
def product_create():
    tenant, schema_name = _product_target_ctx()
    has_products_table = _table_exists(schema_name, "Products")
    has_catalog_table = _table_exists(schema_name, "DocProductCatalog")
    if not has_products_table and not has_catalog_table:
        _flash_products_table_missing(schema_name)
        return redirect(_products_url_for_tenant(tenant))

    product = {}
    choices = _product_form_choices()
    columns = _table_columns(schema_name, "Products") if has_products_table else set()
    if request.method == "POST":
        if not has_products_table:
            flash(f"{schema_name}.Products is not available for manual product creation. Use the Excel/CSV import into DocProductCatalog.", "warning")
            return redirect(url_for("master_data.product_create", tenant_code=tenant["code"]))
        product = dict(request.form)
        try:
            payload = _product_form_payload(request.form, columns)
            insert_cols = list(payload.keys())
            placeholders = ["?" for _ in insert_cols]
            params = [payload[col] for col in insert_cols]
            if "created_at" in columns:
                insert_cols.append("created_at")
                placeholders.append("GETUTCDATE()")
            execute(
                "INSERT INTO {table_name} ({columns}) VALUES ({values})".format(
                    table_name=_qualified(schema_name, "Products"),
                    columns=", ".join(insert_cols),
                    values=", ".join(placeholders),
                ),
                params,
            )
            flash(f'Product "{request.form.get("product_name")}" created.', "success")
            return redirect(_products_url_for_tenant(tenant))
        except Exception as exc:
            message = f"Error creating product: {exc}"
            log_id = _log_master_data_event(
                schema_name,
                "MASTERDATA_PRODUCT_CREATE",
                message,
                error_detail=exc,
                request_payload=dict(request.form),
            )
            _flash_message("danger", message, technical_url=_technical_api_url(log_id) if log_id else None)

    return render_template(
        "master_data/product_form.html",
        product=product,
        mode="create",
        tenant=tenant,
        product_tenant=tenant["code"],
        has_products_table=has_products_table,
        has_catalog_table=has_catalog_table,
        **choices,
    )


@master_data_bp.route("/products/<int:product_id>/edit", methods=["GET", "POST"])
def product_edit(product_id):
    tenant, schema_name = _product_target_ctx()
    if not _table_exists(schema_name, "Products"):
        _flash_products_table_missing(schema_name)
        return redirect(_products_url_for_tenant(tenant))

    columns = _table_columns(schema_name, "Products")
    product = query_one(f"{_product_select_sql(schema_name, columns)} WHERE id = ?", [product_id]) or {}
    if not product:
        flash("Product not found.", "warning")
        return redirect(_products_url_for_tenant(tenant))

    choices = _product_form_choices()
    if request.method == "POST":
        product = dict(product)
        product.update(request.form)
        try:
            payload = _product_form_payload(request.form, columns)
            assignments = [f"{col}=?" for col in payload]
            params = [payload[col] for col in payload]
            if "updated_at" in columns:
                assignments.append("updated_at=GETUTCDATE()")
            params.append(product_id)
            execute(
                "UPDATE {table_name} SET {assignments} WHERE id = ?".format(
                    table_name=_qualified(schema_name, "Products"),
                    assignments=", ".join(assignments),
                ),
                params,
            )
            flash("Product updated.", "success")
            return redirect(_products_url_for_tenant(tenant))
        except Exception as exc:
            message = f"Error updating product: {exc}"
            log_id = _log_master_data_event(
                schema_name,
                "MASTERDATA_PRODUCT_UPDATE",
                message,
                error_detail=exc,
                request_payload=dict(request.form),
            )
            _flash_message("danger", message, technical_url=_technical_api_url(log_id) if log_id else None)

    return render_template(
        "master_data/product_form.html",
        product=product,
        mode="edit",
        tenant=tenant,
        product_tenant=tenant["code"],
        **choices,
    )


@master_data_bp.route("/products/catalog/<int:catalog_id>/edit", methods=["GET", "POST"])
def product_catalog_edit(catalog_id):
    tenant, schema_name = _product_target_ctx()
    if not _table_exists(schema_name, "DocProductCatalog"):
        _flash_products_table_missing(schema_name)
        return redirect(_products_url_for_tenant(tenant))

    columns = _table_columns(schema_name, "DocProductCatalog")
    product = query_one(
        f"SELECT * FROM {_qualified(schema_name, 'DocProductCatalog')} WHERE id = ?",
        [catalog_id],
    ) or {}
    if not product:
        flash("Product catalog row not found.", "warning")
        return redirect(_products_url_for_tenant(tenant))

    if request.method == "POST":
        product = dict(product)
        product.update(request.form)
        try:
            payload = _doc_product_catalog_payload_from_form(request.form, columns)
            assignments = [f"[{col}] = ?" for col in payload]
            params = [payload[col] for col in payload]
            if "updated_at" in columns:
                assignments.append("updated_at = SYSUTCDATETIME()")
            params.append(catalog_id)
            execute(
                "UPDATE {table_name} SET {assignments} WHERE id = ?".format(
                    table_name=_qualified(schema_name, "DocProductCatalog"),
                    assignments=", ".join(assignments),
                ),
                params,
            )
            flash("Product catalog row updated.", "success")
            return redirect(_products_url_for_tenant(tenant))
        except Exception as exc:
            message = f"Error updating product catalog row: {exc}"
            log_id = _log_master_data_event(
                schema_name,
                "MASTERDATA_PRODUCT_CATALOG_UPDATE",
                message,
                error_detail=exc,
                request_payload=dict(request.form),
            )
            _flash_message("danger", message, technical_url=_technical_api_url(log_id) if log_id else None)

    return render_template(
        "master_data/product_catalog_form.html",
        product=product,
        tenant=tenant,
        product_tenant=tenant["code"],
        columns=columns,
        package_types=_load_cv_options("CV_type_of_package"),
        procedure_codes=_load_cv_options("CV_procedure_code"),
        additional_procedure_codes=_load_cv_options("CV_additional_procedure_code"),
        valuation_methods=_load_cv_options("CV_valuation_method"),
        valuation_indicators=_load_cv_options("CV_valuation_indicator"),
        preferences=_load_cv_options("CV_preference"),
        ni_codes=_load_cv_options("CV_ni_additional_information_code"),
        nature_of_transaction_options=_load_cv_options("CV_nature_of_transaction"),
        controlled_goods_types=_load_cv_options("CV_controlled_goods_type"),
    )


@master_data_bp.route("/choice-values")
@master_data_bp.route("/cv-tables")
def cv_tables():
    conn = get_db()
    cursor = conn.cursor()
    tables = []
    for table_name in _cv_table_names():
        try:
            cursor.execute(f"SELECT COUNT(*) AS cnt FROM TSS.[{table_name}]")
            count = int(cursor.fetchone()[0])
        except Exception:
            count = None
        tables.append({"name": table_name, "count": count})

    total_rows = sum(t["count"] for t in tables if t["count"] is not None)
    return render_template("master_data/cv_tables.html", tables=tables, total_rows=total_rows)


@master_data_bp.route("/choice-values/<string:table_name>")
@master_data_bp.route("/cv-tables/<string:table_name>")
def cv_table_detail(table_name):
    allowed_tables = set(_cv_table_names()) | set(STATIC_CV_TABLE_NAMES)
    if table_name not in allowed_tables:
        abort(404)

    try:
        page = max(1, int(request.args.get("page", 1)))
    except (TypeError, ValueError):
        page = 1
    search = (request.args.get("q") or "").strip()

    conn = get_db()
    cursor = conn.cursor()

    try:
        cursor.execute(f"SELECT TOP 0 * FROM TSS.[{table_name}]")
        columns = [col[0] for col in cursor.description] if cursor.description else []
    except Exception:
        abort(404)

    search_cols = [col for col in columns if col.lower() in {"value", "name", "location_code", "operator_facility_name", "description"}]
    if not search_cols:
        search_cols = columns[:3]

    where_sql = ""
    where_params = []
    if search and search_cols:
        where_sql = " WHERE " + " OR ".join(f"CAST([{col}] AS NVARCHAR(500)) LIKE ?" for col in search_cols)
        where_params = [f"%{search}%"] * len(search_cols)

    try:
        cursor.execute(f"SELECT COUNT(*) AS cnt FROM TSS.[{table_name}]{where_sql}", where_params)
        total = int(cursor.fetchone()[0])
    except Exception:
        total = 0

    total_pages = max(1, (total + CV_PAGE_SIZE - 1) // CV_PAGE_SIZE)
    page = min(page, total_pages)
    offset = (page - 1) * CV_PAGE_SIZE

    try:
        cursor.execute(
            f"SELECT * FROM TSS.[{table_name}]{where_sql} ORDER BY 1 OFFSET ? ROWS FETCH NEXT {CV_PAGE_SIZE} ROWS ONLY",
            where_params + [offset],
        )
        fetched = cursor.fetchall()
        rows = [dict(zip(columns, row)) for row in fetched]
    except Exception:
        rows = []

    return render_template(
        "master_data/cv_table_detail.html",
        table_name=table_name,
        columns=columns,
        rows=rows,
        search=search,
        page=page,
        total=total,
        total_pages=total_pages,
    )


@master_data_bp.route("/eori-checker", methods=["GET", "POST"])
def eori_checker():
    raw = ""
    results = []
    error = None

    if request.method == "POST":
        raw = (request.form.get("eoris") or "").strip()
        eoris = []
        seen = set()
        for line in raw.splitlines():
            eori = line.strip().upper()
            if eori and eori not in seen:
                eoris.append(eori)
                seen.add(eori)

        if not eoris:
            error = "No EORIs entered."
        elif len(eoris) > 100:
            error = f"Maximum 100 EORIs per batch. You entered {len(eoris)}."
        else:
            try:
                batch = check_eori_batch(eoris)
                results = [batch.get(eori, {"eori": eori, "route": "UNKNOWN", "valid": False, "status": "Error", "error": "No result"}) for eori in eoris]
            except Exception as exc:
                error = f"Validation error: {exc}"

    return render_template("master_data/eori_checker.html", raw=raw, results=results, error=error)
