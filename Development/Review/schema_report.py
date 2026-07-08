#!/usr/bin/env python3
"""Fusion Flow V3 QAS - database schema report to Excel (rows + columns + script links).

Writes an .xlsx to Documentation/ that lists every table with its row count and column
count, the SQL migration that CREATEs it, and which repo scripts (SQL + Python) reference
it - so tables left over from two solutions can be spotted and cleaned up.

    python Development/Review/schema_report.py
    python Development/Review/schema_report.py --out "C:\\...\\Fusion_Flow_V3_QAS\\Documentation"

Connects via DB_* env vars, else Configuration/Fusion_Flow_QAS.ini. Needs pyodbc + openpyxl.
"""

from __future__ import annotations

import argparse
import configparser
import os
import re
from datetime import datetime
from pathlib import Path

REPO_ROOT = Path(__file__).resolve().parents[2]
DEFAULT_INI = REPO_ROOT / "Configuration" / "Fusion_Flow_QAS.ini"
DEFAULT_OUT = REPO_ROOT / "Documentation"
# Where to look for table references (both live solutions + tooling).
SCAN_DIRS = ["Configuration", "Modules", "Development", "liveWeb",
             "Integration_Layer", "Database_Layer", "Deprecated", "Archive"]
SCAN_EXT = (".sql", ".py")


def _config_from_env():
    if not os.environ.get("DB_SERVER"):
        return None
    return {"server": os.environ.get("DB_SERVER", ""), "database": os.environ.get("DB_NAME", ""),
            "user": os.environ.get("DB_USER", ""), "password": os.environ.get("DB_PASSWORD", ""),
            "driver": os.environ.get("DB_DRIVER", "{ODBC Driver 18 for SQL Server}"),
            "encrypt": os.environ.get("DB_ENCRYPT", "yes"),
            "trust_server_certificate": os.environ.get("DB_TRUST", "no")}


def load_db_config(ini_path: Path) -> dict:
    env = _config_from_env()
    if env:
        return env
    if not Path(ini_path).exists():
        raise SystemExit(f"No DB_SERVER env and no connection file: {ini_path}")
    cp = configparser.ConfigParser(); cp.read(ini_path, encoding="utf-8")
    if "database" not in cp:
        raise SystemExit(f"No [database] section in {ini_path}")
    return {k.lower(): v for k, v in cp["database"].items()}


def conn_str(db: dict) -> str:
    yes = lambda v: str(v).lower() in ("yes", "true", "1")
    p = [f"Driver={db.get('driver', '{ODBC Driver 18 for SQL Server}')}",
         f"Server={db['server']}", f"Database={db['database']}"]
    p += ([f"Uid={db['user']}", f"Pwd={db.get('password', '')}"] if db.get("user")
          else ["Trusted_Connection=yes"])
    p.append(f"Encrypt={'yes' if yes(db.get('encrypt', 'yes')) else 'no'}")
    p.append(f"TrustServerCertificate={'yes' if yes(db.get('trust_server_certificate', 'no')) else 'no'}")
    return ";".join(p) + ";"


def scan_repo():
    files = []
    for d in SCAN_DIRS:
        base = REPO_ROOT / d
        if base.exists():
            files += [p for p in base.rglob("*") if p.suffix.lower() in SCAN_EXT]
    contents = {p: p.read_text(encoding="utf-8", errors="ignore") for p in files}
    return files, contents


def refs_for(schema: str, table: str, files, contents):
    """Files that reference schema.table, and the SQL file that CREATEs it (if any)."""
    ref = re.compile(rf"\b{re.escape(schema)}\.\[?{re.escape(table)}\]?\b", re.I)
    crt = re.compile(rf"CREATE\s+TABLE\s+(\[?{re.escape(schema)}\]?\.)?\[?{re.escape(table)}\]?\b", re.I)
    refs, created = [], None
    for p in files:
        txt = contents[p]
        if ref.search(txt) or crt.search(txt):
            rel = str(p.relative_to(REPO_ROOT)).replace("\\", "/")
            refs.append(rel)
            if created is None and p.suffix.lower() == ".sql" and crt.search(txt):
                created = rel
    return refs, created


def main() -> int:
    ap = argparse.ArgumentParser(description="DB schema -> Excel with row counts + script links.")
    ap.add_argument("--ini", default=str(DEFAULT_INI))
    ap.add_argument("--out", default=str(DEFAULT_OUT), help="output folder (default Documentation/)")
    args = ap.parse_args()

    import pyodbc
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    conn = pyodbc.connect(conn_str(load_db_config(Path(args.ini))), autocommit=True)
    cur = conn.cursor()
    tables = cur.execute(
        "SELECT s.name, t.name, SUM(CASE WHEN p.index_id IN (0,1) THEN p.rows ELSE 0 END) AS rws "
        "FROM sys.tables t JOIN sys.schemas s ON s.schema_id = t.schema_id "
        "LEFT JOIN sys.partitions p ON p.object_id = t.object_id "
        "GROUP BY s.name, t.name ORDER BY s.name, t.name").fetchall()
    cols = cur.execute(
        "SELECT TABLE_SCHEMA, TABLE_NAME, COLUMN_NAME, DATA_TYPE, CHARACTER_MAXIMUM_LENGTH, "
        "IS_NULLABLE, ORDINAL_POSITION FROM INFORMATION_SCHEMA.COLUMNS "
        "ORDER BY TABLE_SCHEMA, TABLE_NAME, ORDINAL_POSITION").fetchall()
    conn.close()

    colcount: dict = {}
    for c in cols:
        colcount[(c[0], c[1])] = colcount.get((c[0], c[1]), 0) + 1

    files, contents = scan_repo()

    HDR = Font(bold=True, color="FFFFFF")
    FILL = PatternFill("solid", fgColor="0F2A4A")
    EMPTY = PatternFill("solid", fgColor="FDE7E7")     # highlight 0-row tables
    ORPHAN = PatternFill("solid", fgColor="FFF3CD")    # highlight tables with no script refs
    wb = Workbook()

    # --- Tables sheet ---
    ws = wb.active; ws.title = "Tables"
    head = ["Schema", "Table", "Rows", "Columns", "Created by (SQL)", "# Script refs", "Referenced in"]
    ws.append(head)
    schema_rows: dict = {}
    for sch, tbl, rws in tables:
        refs, created = refs_for(sch, tbl, files, contents)
        ws.append([sch, tbl, int(rws or 0), colcount.get((sch, tbl), 0),
                   created or "—", len(refs), ", ".join(refs) if refs else "(none)"])
        r = ws.max_row
        if int(rws or 0) == 0:
            ws.cell(r, 3).fill = EMPTY
        if not refs:
            ws.cell(r, 7).fill = ORPHAN
        schema_rows[sch] = schema_rows.get(sch, [0, 0])
        schema_rows[sch][0] += 1
        schema_rows[sch][1] += int(rws or 0)

    # --- Columns sheet ---
    wc = wb.create_sheet("Columns")
    wc.append(["Schema", "Table", "Column", "Type", "Max len", "Nullable", "Ordinal"])
    for c in cols:
        wc.append([c[0], c[1], c[2], c[3], c[4], c[5], c[6]])

    # --- Table <-> Scripts map ---
    wm = wb.create_sheet("Table-Scripts")
    wm.append(["Schema", "Table", "Rows", "Script"])
    for sch, tbl, rws in tables:
        refs, _ = refs_for(sch, tbl, files, contents)
        if refs:
            for f in refs:
                wm.append([sch, tbl, int(rws or 0), f])
        else:
            wm.append([sch, tbl, int(rws or 0), "(no script references — cleanup candidate)"])

    # --- Overview ---
    wo = wb.create_sheet("Overview", 0)
    wo.append(["Fusion Flow V3 QAS — schema report"])
    wo.append([f"Generated {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"])
    wo.append([])
    wo.append(["Schema", "Tables", "Total rows"])
    for sch in sorted(schema_rows):
        wo.append([sch, schema_rows[sch][0], schema_rows[sch][1]])
    wo.append([])
    wo.append(["Legend", "pink = 0 rows (empty)", "amber = no script references (cleanup candidate)"])

    # style headers + widths
    for sheet, hdr_row in ((ws, 1), (wc, 1), (wm, 1), (wo, 4)):
        for cell in sheet[hdr_row]:
            if cell.value:
                cell.font = HDR; cell.fill = FILL; cell.alignment = Alignment(vertical="center")
        sheet.freeze_panes = sheet.cell(hdr_row + 1, 1)
    for sheet in (ws, wc, wm, wo):
        for i, col in enumerate(sheet.columns, 1):
            width = min(90, max(10, max((len(str(c.value)) for c in col if c.value), default=10) + 2))
            sheet.column_dimensions[get_column_letter(i)].width = width
    wo["A1"].font = Font(bold=True, size=14)

    os.makedirs(args.out, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    out = Path(args.out) / f"Schema_Report_{stamp}.xlsx"
    wb.save(out)
    print(f"Wrote {out}")
    print(f"Tables: {len(tables)}  ·  empty: {sum(1 for _,_,r in tables if int(r or 0)==0)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
